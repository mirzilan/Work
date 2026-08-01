from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.workbook import Workbook

import assumptions_constant as const
import calc_capex as capex
import calc_cfads as cfads
import calc_financing_cons as fin_cons
import calc_financing_ops as fin_ops
import calc_tax as tax
from inputs import ProjectInputs
from timeline import Timeline
# Source-sheet row constants, imported rather than restated — see the note in calc_tax.py.
from calc_revenue_opex import (
    ROW_REVENUE as REVOPEX_ROW_REVENUE,
    ROW_OPEX as REVOPEX_ROW_OPEX,
    ROW_EBITDA as REVOPEX_ROW_EBITDA,
)
from workbook_builder import (
    FIRST_DATA_COL,
    COLOR_INPUT,
    COLOR_FORMULA,
    COLOR_LINK,
    TAB_COLOR_OUTPUT,
    col_letter,
    style_total_row,
    style_section_header_row,
)

# Flags block — frozen with the timeline below it, so scenario/status context stays
# visible no matter how far the user scrolls. See Legend for what each flag means.
ROW_FLAG_ACTIVE_SCENARIO = 2
ROW_FLAG_MODEL_STATUS = 3
ROW_FLAG_SOLVE_FRESHNESS = 4

ROW_DATE_HEADER = 7
ROW_QUARTER_INDEX = 8

# ---------------------------------------------------------------------------------
# P&L
# ---------------------------------------------------------------------------------
ROW_PNL_HEADER = 10
ROW_REVENUE = 11
ROW_OPEX = 12
ROW_EBITDA = 13
ROW_DEPRECIATION = 14
ROW_EBIT = 15
ROW_INTEREST_EXPENSE = 16
ROW_EBT = 17
ROW_TAX = 18
ROW_LC_FEE = 19  # below tax: non-deductible, see the note in calc_tax.py
ROW_NET_INCOME = 20

# ---------------------------------------------------------------------------------
# BALANCE SHEET — classified (current/non-current), A = L + E displayed explicitly
# ---------------------------------------------------------------------------------
ROW_BS_HEADER = 22
ROW_BS_ASSETS_HEADER = 23
ROW_BS_CURRENT_ASSETS_HEADER = 24
ROW_BS_CASH = 25
ROW_BS_TOTAL_CURRENT_ASSETS = 26
ROW_BS_NONCURRENT_ASSETS_HEADER = 27
ROW_BS_DSRA = 28
ROW_BS_MRA = 29
ROW_BS_PPE_NET = 30
ROW_BS_TOTAL_NONCURRENT_ASSETS = 31
ROW_BS_TOTAL_ASSETS = 32

ROW_BS_LIABILITIES_HEADER = 34
ROW_BS_CURRENT_LIAB_HEADER = 35
ROW_BS_DEBT_CURRENT = 36  # portion due within the next 4 quarters
ROW_BS_TOTAL_CURRENT_LIAB = 37
ROW_BS_NONCURRENT_LIAB_HEADER = 38
ROW_BS_DEBT_NONCURRENT = 39
ROW_BS_TOTAL_NONCURRENT_LIAB = 40
ROW_BS_TOTAL_LIABILITIES = 41

ROW_BS_EQUITY_HEADER = 43
ROW_BS_PAID_IN_CAPITAL = 44
ROW_BS_RETAINED_EARNINGS = 45
ROW_BS_TOTAL_EQUITY = 46

ROW_BS_TOTAL_LIAB_EQUITY = 48
ROW_BS_CHECK_A_MINUS_L = 49  # displays Assets - Liabilities for a visual E = A - L proof

# ---------------------------------------------------------------------------------
# CASH FLOW STATEMENT — indirect method. DSRA/MRA reserve funding is NOT shown within
# CFF: per ASC 230-10-45 (restricted cash), moving cash into a reserve account is a
# reclassification within "cash and cash equivalents, including restricted cash," not an
# operating/investing/financing activity. The statement below therefore walks to a
# Total Cash figure (unrestricted + DSRA + MRA), then reconciles down to unrestricted
# cash as a schedule beneath it — the "adjustment for restricted cash" sits below the
# main statement rather than being netted into CFF.
# ---------------------------------------------------------------------------------
ROW_CF_HEADER = 51
ROW_CFO_NI = 52
ROW_CFO_ADDBACK_DEPR = 53
ROW_CFO_WC_CHANGE = 54  # placeholder — see note at the cell itself
ROW_CFO = 55
ROW_CFI = 56
ROW_CFF_PRINCIPAL = 57
ROW_CFF_DIVIDENDS = 58
ROW_CFF_EQUITY_INJECTION = 59
ROW_CFF = 60
ROW_NET_CHANGE_TOTAL_CASH = 61
ROW_OPENING_TOTAL_CASH = 62
ROW_CLOSING_TOTAL_CASH = 63

ROW_CF_RESTRICTED_HEADER = 65
ROW_CF_LESS_DSRA = 66
ROW_CF_LESS_MRA = 67
ROW_CLOSING_CASH = 68  # closing UNRESTRICTED cash — this is what feeds the BS cash line

# Direct method — cross-check on CFO only (no AR/AP/inventory anywhere in the model, so
# cash receipts/payments equal the P&L lines exactly; a mismatch means the indirect build
# or a source link is wrong, not a timing difference).
ROW_CFD_HEADER = 70
ROW_CFD_RECEIPTS = 71
ROW_CFD_OPEX_PAID = 72
ROW_CFD_INTEREST_PAID = 73
ROW_CFD_TAX_PAID = 74
ROW_CFD_LC_FEE_PAID = 75
ROW_CFO_DIRECT = 76
ROW_CHECK_DIRECT_TIES_INDIRECT = 77

ROW_CHECK_HEADER = 79
ROW_CHECK_BS_BALANCES_COUNT = 80         # count of quarters where the BS does not balance
ROW_CHECK_CASH_TIES_BUFFER = 81          # closing unrestricted cash ties to Calc_CFADS' buffer
ROW_CHECK_DIRECT_TIES_INDIRECT_COUNT = 82  # count of quarters where direct CFO != indirect CFO

# ---------------------------------------------------------------------------------
# FREE CASH FLOW — two independently-derived views of FCFF and FCFE, cross-checked.
#
# FCFF: "Cash Flow Statement method" (CFO + CFI + Interest + LC Fee — un-levering the
# built statement) vs "CFADS method" (CFADS - Maintenance Capex, linked from Calc_CFADS).
# These are the same figure by algebraic identity (CFADS = EBITDA - Tax = NI + Depr +
# Interest + LC Fee), so the check below expects an EXACT match every quarter — any gap
# means a wiring error, not a timing difference.
#
# FCFE: "Cash Flow Statement method" (CFO + CFI - Principal - DSRA Funding - MRA Funding
# — cash available to equity BEFORE the buffer/lock-up policy decides how much to actually
# pay out) vs "Dividend method" (Distributions - Equity Injections, linked from
# Calc_CFADS — the actual cash equity receives, and what the reported EIRR is built on).
# These two do NOT tie quarter-by-quarter whenever the buffer is building, releasing, or a
# DSCR lock-up blocks a distribution — that's the buffer smoothing timing, not a bug.
#
# Over the whole life they tie to exactly the Initial Cash Buffer funded at Financial
# Close, not to zero: that buffer is funded once, outside CAFD entirely (it comes from the
# construction-period Sources & Uses, not from any operating quarter's cash flow), then
# drains out as part of the distribution once the buffer target is forced to zero in the
# final quarter. So SUM(Distributions - Injections) = SUM(CAFD) + Initial Buffer — proven
# from the buffer roll-forward (Distribution - Injection = CAFD - Change in Buffer, summed
# across all quarters, where the buffer starts at its Financial-Close value and ends at 0).
# ---------------------------------------------------------------------------------
ROW_FCF_HEADER = 84
ROW_FCFF_CF_METHOD = 85
ROW_FCFF_CFADS_METHOD = 86
ROW_CHECK_FCFF_METHODS_TIE_COUNT = 87

ROW_FCFE_CF_METHOD = 89
ROW_FCFE_DIVIDEND_METHOD = 90
ROW_CHECK_FCFE_LIFETIME_TIE = 91


def build_fs_quarterly(wb: Workbook, timeline: Timeline, inputs: ProjectInputs) -> Worksheet:
    ws = wb.create_sheet("FS_Quarterly")
    ws.sheet_properties.tabColor = TAB_COLOR_OUTPUT

    ws["A1"] = "FS_Quarterly — P&L / Balance Sheet / Cash Flow (reserves as restricted cash, 100% FCFE payout)"
    ws["A1"].font = Font(bold=True, size=12)

    _build_flags_block(ws)

    _label(ws, ROW_DATE_HEADER, "Period End Date")
    _label(ws, ROW_QUARTER_INDEX, "Operating Quarter #")

    _section_header(ws, ROW_PNL_HEADER, "PROFIT & LOSS")
    _label(ws, ROW_REVENUE, "Revenue ($)")
    _label(ws, ROW_OPEX, "Opex ($)")
    _label(ws, ROW_EBITDA, "EBITDA ($)")
    _label(ws, ROW_DEPRECIATION, "Depreciation ($) — base vintage + maintenance vintages")
    _label(ws, ROW_EBIT, "EBIT ($)")
    _label(ws, ROW_INTEREST_EXPENSE, "Interest Expense ($)")
    _label(ws, ROW_EBT, "EBT ($)")
    _label(ws, ROW_TAX, "Tax ($) — linked from Calc_Tax")
    _label(ws, ROW_LC_FEE, "DSRA LC Fee ($) — non-deductible; zero when the DSRA is cash-funded")
    _bold_label(ws, ROW_NET_INCOME, "Net Income ($)")

    _section_header(ws, ROW_BS_HEADER, "BALANCE SHEET")
    _bold_label(ws, ROW_BS_ASSETS_HEADER, "ASSETS")
    _sub_label(ws, ROW_BS_CURRENT_ASSETS_HEADER, "Current Assets")
    _label(ws, ROW_BS_CASH, "Cash & Cash Equivalents ($) — unrestricted")
    _bold_label(ws, ROW_BS_TOTAL_CURRENT_ASSETS, "Total Current Assets ($)")
    _sub_label(ws, ROW_BS_NONCURRENT_ASSETS_HEADER, "Non-Current Assets")
    _label(ws, ROW_BS_DSRA, "DSRA Balance ($) — restricted cash")
    _label(ws, ROW_BS_MRA, "MRA Balance ($) — restricted cash")
    _label(ws, ROW_BS_PPE_NET, "PP&E, Net ($)")
    _bold_label(ws, ROW_BS_TOTAL_NONCURRENT_ASSETS, "Total Non-Current Assets ($)")
    _bold_label(ws, ROW_BS_TOTAL_ASSETS, "TOTAL ASSETS ($)")

    _bold_label(ws, ROW_BS_LIABILITIES_HEADER, "LIABILITIES")
    _sub_label(ws, ROW_BS_CURRENT_LIAB_HEADER, "Current Liabilities")
    _label(ws, ROW_BS_DEBT_CURRENT, "Debt — Current Portion ($) — due within 4 quarters")
    _bold_label(ws, ROW_BS_TOTAL_CURRENT_LIAB, "Total Current Liabilities ($)")
    _sub_label(ws, ROW_BS_NONCURRENT_LIAB_HEADER, "Non-Current Liabilities")
    _label(ws, ROW_BS_DEBT_NONCURRENT, "Debt — Non-Current Portion ($)")
    _bold_label(ws, ROW_BS_TOTAL_NONCURRENT_LIAB, "Total Non-Current Liabilities ($)")
    _bold_label(ws, ROW_BS_TOTAL_LIABILITIES, "TOTAL LIABILITIES ($)")

    _bold_label(ws, ROW_BS_EQUITY_HEADER, "EQUITY")
    _label(ws, ROW_BS_PAID_IN_CAPITAL, "Paid-in Capital ($)")
    _label(ws, ROW_BS_RETAINED_EARNINGS, "Retained Earnings ($)")
    _bold_label(ws, ROW_BS_TOTAL_EQUITY, "TOTAL EQUITY ($)")

    _bold_label(ws, ROW_BS_TOTAL_LIAB_EQUITY, "TOTAL LIABILITIES + EQUITY ($)")
    _label(ws, ROW_BS_CHECK_A_MINUS_L, "Check: Assets − Liabilities (should equal Total Equity above)")

    _section_header(ws, ROW_CF_HEADER, "CASH FLOW STATEMENT — INDIRECT METHOD")
    _label(ws, ROW_CFO_NI, "Net Income ($)")
    _label(ws, ROW_CFO_ADDBACK_DEPR, "Add back: Depreciation ($)")
    _label(ws, ROW_CFO_WC_CHANGE,
           "Change in Working Capital ($) — placeholder; no AR/AP/inventory modelled (Phase 2/3 scope)")
    _bold_label(ws, ROW_CFO, "Cash Flow from Operations ($)")
    _bold_label(ws, ROW_CFI, "Cash Flow from Investing ($) — maintenance capex")
    _label(ws, ROW_CFF_PRINCIPAL, "Debt Principal Repayment ($)")
    _label(ws, ROW_CFF_DIVIDENDS, "Distributions to Equity ($)")
    _label(ws, ROW_CFF_EQUITY_INJECTION, "Equity Injections ($) — shortfalls the buffer could not cover")
    _bold_label(ws, ROW_CFF, "Cash Flow from Financing ($)")
    _bold_label(ws, ROW_NET_CHANGE_TOTAL_CASH,
                "Net Change in Total Cash ($) — incl. restricted; DSRA/MRA funding is a reclass, not a cash flow")
    _label(ws, ROW_OPENING_TOTAL_CASH, "Opening Total Cash ($) — unrestricted + DSRA + MRA")
    _bold_label(ws, ROW_CLOSING_TOTAL_CASH, "Closing Total Cash ($) — unrestricted + DSRA + MRA")

    _section_header(ws, ROW_CF_RESTRICTED_HEADER, "Reconciliation — Total Cash to Unrestricted Cash")
    _label(ws, ROW_CF_LESS_DSRA, "Less: DSRA Balance ($)")
    _label(ws, ROW_CF_LESS_MRA, "Less: MRA Balance ($)")
    _bold_label(ws, ROW_CLOSING_CASH, "Closing Cash & Cash Equivalents ($) — unrestricted")

    _section_header(ws, ROW_CFD_HEADER, "CASH FLOW STATEMENT — DIRECT METHOD (cross-check on CFO)")
    _label(ws, ROW_CFD_RECEIPTS, "Cash Received from Customers ($)")
    _label(ws, ROW_CFD_OPEX_PAID, "Cash Paid for Opex ($)")
    _label(ws, ROW_CFD_INTEREST_PAID, "Cash Paid for Interest ($)")
    _label(ws, ROW_CFD_TAX_PAID, "Cash Paid for Tax ($)")
    _label(ws, ROW_CFD_LC_FEE_PAID, "Cash Paid — DSRA LC Fee ($)")
    _bold_label(ws, ROW_CFO_DIRECT, "Cash Flow from Operations ($) — Direct Method")
    _label(ws, ROW_CHECK_DIRECT_TIES_INDIRECT, "Check: Direct CFO = Indirect CFO")

    ws.cell(row=ROW_CHECK_HEADER, column=1, value="Checks").font = Font(bold=True)
    _label(ws, ROW_CHECK_BS_BALANCES_COUNT, "# of quarters where BS does not balance")
    _label(ws, ROW_CHECK_CASH_TIES_BUFFER, "Check: Closing cash ties to the Calc_CFADS buffer")
    _label(ws, ROW_CHECK_DIRECT_TIES_INDIRECT_COUNT, "# of quarters where Direct CFO != Indirect CFO")

    _section_header(ws, ROW_FCF_HEADER, "FREE CASH FLOW — TWO VIEWS, CROSS-CHECKED")
    _label(ws, ROW_FCFF_CF_METHOD, "FCFF ($) — Cash Flow Statement method = CFO + CFI + Interest + LC Fee")
    _label(ws, ROW_FCFF_CFADS_METHOD,
           "FCFF ($) — CFADS method (Calc_CFADS); used for reported PIRR")
    _label(ws, ROW_CHECK_FCFF_METHODS_TIE_COUNT, "# of quarters where the two FCFF views differ (should be 0)")
    _label(ws, ROW_FCFE_CF_METHOD,
           "FCFE ($) — Cash Flow Statement method = CFO + CFI - Principal - DSRA Funding - MRA Funding "
           "(pre-distribution-policy)")
    _label(ws, ROW_FCFE_DIVIDEND_METHOD,
           "FCFE ($) — Dividend method (Calc_CFADS); actual cash to/from equity, used for reported EIRR")
    _label(ws, ROW_CHECK_FCFE_LIFETIME_TIE,
           "Check: whole-of-life FCFE (Dividend Method) = whole-of-life FCFE (CF Method) + Initial Cash "
           "Buffer — quarterly timing differs by design (buffer/lock-up); this identity must not")

    n_quarters = len(timeline.operations_quarters)
    last_cons_col = col_letter(len(timeline.construction_months) - 1)
    last_q_col = col_letter(n_quarters - 1)

    for i, period in enumerate(timeline.operations_quarters):
        col = col_letter(i)
        prev_col = col_letter(i - 1) if i > 0 else None

        ws[f"{col}{ROW_DATE_HEADER}"] = period.end
        ws[f"{col}{ROW_DATE_HEADER}"].number_format = "mmm-yy"
        ws[f"{col}{ROW_QUARTER_INDEX}"] = i + 1

        # --- P&L ---
        _link(ws, col, ROW_REVENUE, f"Calc_Revenue_Opex!{col}{REVOPEX_ROW_REVENUE}")
        _link(ws, col, ROW_OPEX, f"Calc_Revenue_Opex!{col}{REVOPEX_ROW_OPEX}")
        _link(ws, col, ROW_EBITDA, f"Calc_Revenue_Opex!{col}{REVOPEX_ROW_EBITDA}")
        _link(ws, col, ROW_DEPRECIATION, f"Calc_Tax!{col}{tax.ROW_TOTAL_DEPRECIATION}")
        _formula(ws, col, ROW_EBIT, f"={col}{ROW_EBITDA}-{col}{ROW_DEPRECIATION}")
        _link(ws, col, ROW_INTEREST_EXPENSE, f"Calc_Financing_Ops!{col}{fin_ops.ROW_INTEREST}")
        _formula(ws, col, ROW_EBT, f"={col}{ROW_EBIT}-{col}{ROW_INTEREST_EXPENSE}")
        _link(ws, col, ROW_TAX, f"Calc_Tax!{col}{tax.ROW_TAX}")
        _link(ws, col, ROW_LC_FEE, f"Calc_Financing_Ops!{col}{fin_ops.ROW_DSRA_LC_FEE}")
        _formula(ws, col, ROW_NET_INCOME, f"={col}{ROW_EBT}-{col}{ROW_TAX}-{col}{ROW_LC_FEE}", bold=True)

        # --- Cash Flow (computed before the Balance Sheet cells that reference it — Excel
        # resolves by dependency, not by row order, so the BS block above can safely point
        # at ROW_CLOSING_CASH even though it's laid out further down the sheet) ---
        _formula(ws, col, ROW_CFO_NI, f"={col}{ROW_NET_INCOME}")
        _formula(ws, col, ROW_CFO_ADDBACK_DEPR, f"={col}{ROW_DEPRECIATION}")
        wc_cell = ws[f"{col}{ROW_CFO_WC_CHANGE}"]
        wc_cell.value = 0.0
        wc_cell.font = Font(color=COLOR_INPUT)
        wc_cell.number_format = "#,##0"
        _formula(ws, col, ROW_CFO,
                 f"={col}{ROW_CFO_NI}+{col}{ROW_CFO_ADDBACK_DEPR}+{col}{ROW_CFO_WC_CHANGE}", bold=True)
        _formula(ws, col, ROW_CFI, f"=-Calc_CFADS!{col}{cfads.ROW_MAINT_CAPEX}", bold=True)
        _link(ws, col, ROW_CFF_PRINCIPAL, f"Calc_Financing_Ops!{col}{fin_ops.ROW_PRINCIPAL}")
        _link(ws, col, ROW_CFF_DIVIDENDS, f"Calc_CFADS!{col}{cfads.ROW_DISTRIBUTION}")
        _link(ws, col, ROW_CFF_EQUITY_INJECTION, f"Calc_CFADS!{col}{cfads.ROW_EQUITY_INJECTION}")
        _formula(ws, col, ROW_CFF,
                 f"=-{col}{ROW_CFF_PRINCIPAL}-{col}{ROW_CFF_DIVIDENDS}+{col}{ROW_CFF_EQUITY_INJECTION}", bold=True)
        _formula(ws, col, ROW_NET_CHANGE_TOTAL_CASH,
                 f"={col}{ROW_CFO}+{col}{ROW_CFI}+{col}{ROW_CFF}", bold=True)

        if i == 0:
            _formula(ws, col, ROW_OPENING_TOTAL_CASH,
                     f"=Calc_Financing_Cons!{fin_cons.ABS_INITIAL_BUFFER}"
                     f"+Calc_Financing_Cons!{fin_cons.ABS_INITIAL_DSRA}")
        else:
            _formula(ws, col, ROW_OPENING_TOTAL_CASH, f"={prev_col}{ROW_CLOSING_TOTAL_CASH}")
        _formula(ws, col, ROW_CLOSING_TOTAL_CASH,
                 f"={col}{ROW_OPENING_TOTAL_CASH}+{col}{ROW_NET_CHANGE_TOTAL_CASH}", bold=True)

        # Reserves are linked here first since the reconciliation and the BS both need them.
        _link(ws, col, ROW_BS_DSRA, f"Calc_Financing_Ops!{col}{fin_ops.ROW_DSRA_BALANCE}")
        _link(ws, col, ROW_BS_MRA, f"Calc_CFADS!{col}{cfads.ROW_MRA_BALANCE}")

        _formula(ws, col, ROW_CF_LESS_DSRA, f"=-{col}{ROW_BS_DSRA}")
        _formula(ws, col, ROW_CF_LESS_MRA, f"=-{col}{ROW_BS_MRA}")
        _formula(ws, col, ROW_CLOSING_CASH,
                 f"={col}{ROW_CLOSING_TOTAL_CASH}+{col}{ROW_CF_LESS_DSRA}+{col}{ROW_CF_LESS_MRA}", bold=True)

        # --- Direct method cross-check ---
        _formula(ws, col, ROW_CFD_RECEIPTS, f"={col}{ROW_REVENUE}")
        _formula(ws, col, ROW_CFD_OPEX_PAID, f"=-{col}{ROW_OPEX}")
        _formula(ws, col, ROW_CFD_INTEREST_PAID, f"=-{col}{ROW_INTEREST_EXPENSE}")
        _formula(ws, col, ROW_CFD_TAX_PAID, f"=-{col}{ROW_TAX}")
        _formula(ws, col, ROW_CFD_LC_FEE_PAID, f"=-{col}{ROW_LC_FEE}")
        _formula(ws, col, ROW_CFO_DIRECT,
                 f"={col}{ROW_CFD_RECEIPTS}+{col}{ROW_CFD_OPEX_PAID}+{col}{ROW_CFD_INTEREST_PAID}"
                 f"+{col}{ROW_CFD_TAX_PAID}+{col}{ROW_CFD_LC_FEE_PAID}", bold=True)
        _formula(ws, col, ROW_CHECK_DIRECT_TIES_INDIRECT,
                 f"=IF(ROUND({col}{ROW_CFO_DIRECT}-{col}{ROW_CFO},2)=0,1,0)")

        # --- Balance Sheet ---
        _formula(ws, col, ROW_BS_CASH, f"={col}{ROW_CLOSING_CASH}")
        _formula(ws, col, ROW_BS_TOTAL_CURRENT_ASSETS, f"={col}{ROW_BS_CASH}", bold=True)

        # CFI is already negative, so subtracting it capitalises the maintenance spend.
        if i == 0:
            _formula(ws, col, ROW_BS_PPE_NET,
                     f"=Calc_Capex!${last_cons_col}${capex.ROW_TOTAL_PROJECT_COST}"
                     f"-{col}{ROW_CFI}-{col}{ROW_DEPRECIATION}")
        else:
            _formula(ws, col, ROW_BS_PPE_NET,
                     f"={prev_col}{ROW_BS_PPE_NET}-{col}{ROW_CFI}-{col}{ROW_DEPRECIATION}")
        _formula(ws, col, ROW_BS_TOTAL_NONCURRENT_ASSETS,
                 f"={col}{ROW_BS_DSRA}+{col}{ROW_BS_MRA}+{col}{ROW_BS_PPE_NET}", bold=True)
        _formula(ws, col, ROW_BS_TOTAL_ASSETS,
                 f"={col}{ROW_BS_TOTAL_CURRENT_ASSETS}+{col}{ROW_BS_TOTAL_NONCURRENT_ASSETS}", bold=True)

        closing_bal_formula = f"Calc_Financing_Ops!{col}{fin_ops.ROW_CLOSING_BAL}"
        current_debt_formula = _current_debt_formula(i, n_quarters, col, closing_bal_formula)
        _formula(ws, col, ROW_BS_DEBT_CURRENT, current_debt_formula)
        _formula(ws, col, ROW_BS_TOTAL_CURRENT_LIAB, f"={col}{ROW_BS_DEBT_CURRENT}", bold=True)
        _formula(ws, col, ROW_BS_DEBT_NONCURRENT, f"={closing_bal_formula}-{col}{ROW_BS_DEBT_CURRENT}")
        _formula(ws, col, ROW_BS_TOTAL_NONCURRENT_LIAB, f"={col}{ROW_BS_DEBT_NONCURRENT}", bold=True)
        _formula(ws, col, ROW_BS_TOTAL_LIABILITIES,
                 f"={col}{ROW_BS_TOTAL_CURRENT_LIAB}+{col}{ROW_BS_TOTAL_NONCURRENT_LIAB}", bold=True)

        if i == 0:
            _formula(ws, col, ROW_BS_PAID_IN_CAPITAL,
                     f"=Calc_Capex!${last_cons_col}${capex.ROW_CUM_EQUITY_DRAW}"
                     f"+{col}{ROW_CFF_EQUITY_INJECTION}")
        else:
            _formula(ws, col, ROW_BS_PAID_IN_CAPITAL,
                     f"={prev_col}{ROW_BS_PAID_IN_CAPITAL}+{col}{ROW_CFF_EQUITY_INJECTION}")
        if i == 0:
            _formula(ws, col, ROW_BS_RETAINED_EARNINGS,
                     f"={col}{ROW_NET_INCOME}-{col}{ROW_CFF_DIVIDENDS}")
        else:
            _formula(ws, col, ROW_BS_RETAINED_EARNINGS,
                     f"={prev_col}{ROW_BS_RETAINED_EARNINGS}+{col}{ROW_NET_INCOME}-{col}{ROW_CFF_DIVIDENDS}")
        _formula(ws, col, ROW_BS_TOTAL_EQUITY,
                 f"={col}{ROW_BS_PAID_IN_CAPITAL}+{col}{ROW_BS_RETAINED_EARNINGS}", bold=True)

        _formula(ws, col, ROW_BS_TOTAL_LIAB_EQUITY,
                 f"={col}{ROW_BS_TOTAL_LIABILITIES}+{col}{ROW_BS_TOTAL_EQUITY}", bold=True)
        _formula(ws, col, ROW_BS_CHECK_A_MINUS_L,
                 f"={col}{ROW_BS_TOTAL_ASSETS}-{col}{ROW_BS_TOTAL_LIABILITIES}")

        # --- FCFF / FCFE: two views ---
        _formula(ws, col, ROW_FCFF_CF_METHOD,
                 f"={col}{ROW_CFO}+{col}{ROW_CFI}+{col}{ROW_INTEREST_EXPENSE}+{col}{ROW_LC_FEE}")
        _link(ws, col, ROW_FCFF_CFADS_METHOD, f"Calc_CFADS!{col}{cfads.ROW_FCFF}")
        _formula(ws, col, ROW_FCFE_CF_METHOD,
                 f"={col}{ROW_CFO}+{col}{ROW_CFI}-{col}{ROW_CFF_PRINCIPAL}"
                 f"-Calc_Financing_Ops!{col}{fin_ops.ROW_DSRA_FUNDING}-Calc_CFADS!{col}{cfads.ROW_MRA_FUNDING}")
        _link(ws, col, ROW_FCFE_DIVIDEND_METHOD, f"Calc_CFADS!{col}{cfads.ROW_FCFE}")

    first_col_idx = FIRST_DATA_COL
    last_col_idx = FIRST_DATA_COL + n_quarters - 1

    # FAST/Corality subtotal convention: single rule above an ordinary subtotal, double
    # rule below a statement's "final answer" (grand=True).
    for row in (
        ROW_BS_TOTAL_CURRENT_ASSETS, ROW_BS_TOTAL_NONCURRENT_ASSETS,
        ROW_BS_TOTAL_CURRENT_LIAB, ROW_BS_TOTAL_NONCURRENT_LIAB, ROW_BS_TOTAL_LIABILITIES,
        ROW_BS_TOTAL_EQUITY, ROW_CFO, ROW_CFI, ROW_CFF, ROW_NET_CHANGE_TOTAL_CASH,
        ROW_CLOSING_TOTAL_CASH, ROW_CFO_DIRECT,
    ):
        style_total_row(ws, row, first_col_idx, last_col_idx)
    for row in (ROW_NET_INCOME, ROW_BS_TOTAL_ASSETS, ROW_BS_TOTAL_LIAB_EQUITY, ROW_CLOSING_CASH):
        style_total_row(ws, row, first_col_idx, last_col_idx, grand=True)

    for row in (ROW_PNL_HEADER, ROW_BS_HEADER, ROW_CF_HEADER, ROW_CF_RESTRICTED_HEADER,
               ROW_CFD_HEADER, ROW_FCF_HEADER):
        style_section_header_row(ws, row, first_col_idx, last_col_idx)

    first_col = col_letter(0)

    check_cell = ws[f"{last_q_col}{ROW_CHECK_BS_BALANCES_COUNT}"]
    check_cell.value = (
        f"=SUMPRODUCT(--(ROUND({first_col}{ROW_BS_TOTAL_ASSETS}:{last_q_col}{ROW_BS_TOTAL_ASSETS}"
        f"-{first_col}{ROW_BS_TOTAL_LIAB_EQUITY}:{last_q_col}{ROW_BS_TOTAL_LIAB_EQUITY},2)<>0))"
    )
    check_cell.font = Font(color=COLOR_FORMULA)

    # An independent tie-out: the statements build cash from CFO/CFI/CFF, while Calc_CFADS
    # builds the same balance from the waterfall. They are separate derivations, so a
    # mismatch means one of them is wrong.
    cash_tie = ws[f"{last_q_col}{ROW_CHECK_CASH_TIES_BUFFER}"]
    cash_tie.value = (
        f"=IF(SUMPRODUCT(--(ROUND({first_col}{ROW_CLOSING_CASH}:{last_q_col}{ROW_CLOSING_CASH}"
        f"-Calc_CFADS!{first_col}{cfads.ROW_BUFFER_CLOSING}:"
        f"Calc_CFADS!{last_q_col}{cfads.ROW_BUFFER_CLOSING},2)<>0))=0,1,0)"
    )
    cash_tie.font = Font(color=COLOR_FORMULA)

    direct_tie_count = ws[f"{last_q_col}{ROW_CHECK_DIRECT_TIES_INDIRECT_COUNT}"]
    direct_tie_count.value = (
        f"=SUMPRODUCT(--({first_col}{ROW_CHECK_DIRECT_TIES_INDIRECT}:"
        f"{last_q_col}{ROW_CHECK_DIRECT_TIES_INDIRECT}=0))"
    )
    direct_tie_count.font = Font(color=COLOR_FORMULA)

    fcff_tie = ws[f"{last_q_col}{ROW_CHECK_FCFF_METHODS_TIE_COUNT}"]
    fcff_tie.value = (
        f"=SUMPRODUCT(--(ROUND({first_col}{ROW_FCFF_CF_METHOD}:{last_q_col}{ROW_FCFF_CF_METHOD}"
        f"-{first_col}{ROW_FCFF_CFADS_METHOD}:{last_q_col}{ROW_FCFF_CFADS_METHOD},2)<>0))"
    )
    fcff_tie.font = Font(color=COLOR_FORMULA)

    # SUM(Distributions - Injections) = SUM(CAFD) + Initial Buffer — see the note at
    # ROW_FCFE_CF_METHOD above for the derivation from the buffer roll-forward.
    fcfe_lifetime_tie = ws[f"{last_q_col}{ROW_CHECK_FCFE_LIFETIME_TIE}"]
    fcfe_lifetime_tie.value = (
        f"=IF(ROUND(SUM({first_col}{ROW_FCFE_DIVIDEND_METHOD}:{last_q_col}{ROW_FCFE_DIVIDEND_METHOD})"
        f"-SUM({first_col}{ROW_FCFE_CF_METHOD}:{last_q_col}{ROW_FCFE_CF_METHOD})"
        f"-Calc_Financing_Cons!{fin_cons.ABS_INITIAL_BUFFER},2)=0,1,0)"
    )
    fcfe_lifetime_tie.font = Font(color=COLOR_FORMULA)

    _add_named_range(wb, "FSQ_LastCol", "FS_Quarterly", f"{last_q_col}1")
    _add_named_range(wb, "FSQ_CashTiesBufferCheck", "FS_Quarterly", f"{last_q_col}{ROW_CHECK_CASH_TIES_BUFFER}")
    _add_named_range(wb, "FSQ_BSBalancesFailCount", "FS_Quarterly", f"{last_q_col}{ROW_CHECK_BS_BALANCES_COUNT}")
    _add_named_range(wb, "FSQ_DirectTiesIndirectFailCount", "FS_Quarterly",
                     f"{last_q_col}{ROW_CHECK_DIRECT_TIES_INDIRECT_COUNT}")

    ws.freeze_panes = ws.cell(row=ROW_PNL_HEADER, column=FIRST_DATA_COL)
    ws.column_dimensions["A"].width = 62

    return ws


def _current_debt_formula(i: int, n_quarters: int, col: str, closing_bal_formula: str) -> str:
    """Debt due within the next 4 quarters, capped at the outstanding balance so it can
    never overstate the current portion near maturity. At the final modelled quarter there
    is no "next 4 quarters" column to sum, so the whole remaining balance (which the
    Calc_Financing_Ops check already proves is ~0 by then) is the current portion."""
    if i == n_quarters - 1:
        return f"={closing_bal_formula}"
    start_idx = i + 1
    end_idx = min(i + 4, n_quarters - 1)
    start_col = col_letter(start_idx)
    end_col = col_letter(end_idx)
    return (
        f"=MIN(SUM(Calc_Financing_Ops!{start_col}{fin_ops.ROW_PRINCIPAL}:"
        f"Calc_Financing_Ops!{end_col}{fin_ops.ROW_PRINCIPAL}),{closing_bal_formula})"
    )


def _build_flags_block(ws: Worksheet) -> None:
    """Active Scenario / Model Status / Solve Freshness — same three flags as Dashboard,
    repeated here since this sheet is read on its own often enough that the reader
    shouldn't have to tab back to Dashboard to know which scenario or state they're in."""
    # Local import: Check_Control aggregates via `import fs_quarterly`, so importing it
    # back at module level here would close the loop — same pattern cover.py already uses.
    import check_control

    _label(ws, ROW_FLAG_ACTIVE_SCENARIO, "Active Scenario")
    scenario_cell = ws[f"B{ROW_FLAG_ACTIVE_SCENARIO}"]
    scenario_cell.value = f"=ActiveScenario&\" - \"&Assumptions_Constant!$B${const.ROW_SCENARIO_NAME}"
    scenario_cell.font = Font(color=COLOR_LINK, bold=True)

    _label(ws, ROW_FLAG_MODEL_STATUS, "Model Status")
    status_cell = ws[f"B{ROW_FLAG_MODEL_STATUS}"]
    status_cell.value = f"=Check_Control!B{check_control.ROW_MASTER_FLAG}"
    status_cell.font = Font(color=COLOR_LINK, bold=True)

    _label(ws, ROW_FLAG_SOLVE_FRESHNESS, "Solve Freshness")
    freshness_cell = ws[f"B{ROW_FLAG_SOLVE_FRESHNESS}"]
    freshness_cell.value = "=SolveStatus"
    freshness_cell.font = Font(color=COLOR_LINK, bold=True)


def _label(ws: Worksheet, row: int, label: str) -> None:
    ws.cell(row=row, column=1, value=label)


def _bold_label(ws: Worksheet, row: int, label: str) -> None:
    ws.cell(row=row, column=1, value=label).font = Font(bold=True)


def _sub_label(ws: Worksheet, row: int, label: str) -> None:
    ws.cell(row=row, column=1, value=label).font = Font(italic=True)


def _section_header(ws: Worksheet, row: int, title: str) -> None:
    cell = ws.cell(row=row, column=1, value=title)
    cell.font = Font(bold=True, size=11, underline="single")


def _link(ws: Worksheet, col: str, row: int, formula: str) -> None:
    cell = ws[f"{col}{row}"]
    cell.value = f"={formula}"
    cell.font = Font(color=COLOR_LINK)
    cell.number_format = "#,##0"


def _formula(ws: Worksheet, col: str, row: int, formula: str, bold: bool = False) -> None:
    cell = ws[f"{col}{row}"]
    cell.value = formula
    cell.font = Font(color=COLOR_FORMULA, bold=bold)
    cell.number_format = "#,##0"


def _add_named_range(wb: Workbook, name: str, sheet: str, cell: str) -> None:
    import re

    from openpyxl.workbook.defined_name import DefinedName

    match = re.match(r"([A-Za-z]+)(\d+)", cell)
    col, row = match.group(1), match.group(2)
    wb.defined_names.add(DefinedName(name, attr_text=f"'{sheet}'!${col}${row}"))
