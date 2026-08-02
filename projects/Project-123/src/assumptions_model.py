from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.workbook import Workbook

from inputs import ProjectInputs
from timeline import Timeline
import assumptions_constant as const
import assumptions_periodic_capex as per_capex
from workbook_builder import FIRST_DATA_COL, COLOR_LINK, TAB_COLOR_INPUT, col_letter

ROW_TOTAL_CAPEX = 3
ROW_DEBT_PCT = 4
ROW_INTEREST_RATE = 5
ROW_DEBT_TENOR_YEARS = 6
ROW_TARGET_DSCR = 7
ROW_ANNUAL_REVENUE = 8
ROW_OPEX_PCT = 9
ROW_TAX_RATE = 10
ROW_USEFUL_LIFE_YEARS = 11

ROW_MAX_GEARING = 12
ROW_ROUTINE_MAINT_PCT = 13
ROW_DSRA_LC_FEE = 14

ROW_PHASING_LABEL = 16
ROW_PHASING_DATE = 17
ROW_PHASING_PCT = 18

# Cell refs (single-scenario, Stage 1b interim — becomes Assumptions_Constant's Active
# column and Assumptions_Periodic_Capex's Active row once Stage 1c scales to 10 scenarios)
CELL_TOTAL_CAPEX = "B3"
CELL_DEBT_PCT = "B4"
CELL_INTEREST_RATE = "B5"
CELL_DEBT_TENOR_YEARS = "B6"
CELL_TARGET_DSCR = "B7"
CELL_ANNUAL_REVENUE = "B8"
CELL_OPEX_PCT = "B9"
CELL_TAX_RATE = "B10"
CELL_USEFUL_LIFE_YEARS = "B11"


def build_assumptions_model(wb: Workbook, timeline: Timeline, inputs: ProjectInputs) -> Worksheet:
    ws = wb.create_sheet("Assumptions_Model")
    ws.sheet_properties.tabColor = TAB_COLOR_INPUT

    ws["A1"] = "Assumptions_Model — Resolved Inputs for the Active Scenario"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = (
        "Every value here resolves from Assumptions_Constant's Active column. Change "
        "assumptions there, not here — this sheet is the live view the Calc sheets read."
    )
    ws["A2"].font = Font(italic=True, size=9)

    # Source rows come from Assumptions_Constant's own constants, so adding a driver there
    # can't silently shift what this sheet resolves.
    _resolved(ws, "A3", "Total Capex ($)", CELL_TOTAL_CAPEX, const.ROW_TOTAL_CAPEX, "#,##0")
    _resolved(ws, "A4", "Gearing (Debt % of TPC)", CELL_DEBT_PCT, const.ROW_GEARING, "0.00%")
    _resolved(ws, "A5", "Interest Rate (Annual)", CELL_INTEREST_RATE, const.ROW_INTEREST_RATE, "0.00%")
    _resolved(ws, "A6", "Debt Tenor (Years)", CELL_DEBT_TENOR_YEARS, const.ROW_DEBT_TENOR, "0")
    _resolved(ws, "A7", "Target DSCR", CELL_TARGET_DSCR, const.ROW_TARGET_DSCR, "0.00x")
    _resolved(ws, "A8", "Annual Revenue ($) — base, pre-index and pre-escalation",
              CELL_ANNUAL_REVENUE, const.ROW_ANNUAL_REVENUE, "#,##0")
    _resolved(ws, "A9", "Opex % of Revenue", CELL_OPEX_PCT, const.ROW_OPEX_PCT, "0.00%")
    _resolved(ws, "A10", "Tax Rate", CELL_TAX_RATE, const.ROW_TAX_RATE, "0.00%")
    _resolved(ws, "A11", "Useful Life (Years)", CELL_USEFUL_LIFE_YEARS, const.ROW_USEFUL_LIFE, "0")
    _resolved(ws, "A12", "Max Gearing (cap on DSCR-sculpted debt size)", "B12", const.ROW_MAX_GEARING, "0.00%")
    _resolved(ws, f"A{ROW_ROUTINE_MAINT_PCT}", "Routine Maint Capex (% of Revenue)",
              f"B{ROW_ROUTINE_MAINT_PCT}", const.ROW_ROUTINE_MAINT_PCT, "0.00%")
    _resolved(ws, f"A{ROW_DSRA_LC_FEE}", "DSRA LC Fee (% p.a. on requirement)",
              f"B{ROW_DSRA_LC_FEE}", const.ROW_DSRA_LC_FEE, "0.00%")

    ws.cell(row=ROW_PHASING_LABEL, column=1,
            value="Capex Phasing % — resolved from Assumptions_Periodic_Capex Active row").font = Font(bold=True)
    ws.cell(row=ROW_PHASING_DATE, column=1, value="Period End Date")
    ws.cell(row=ROW_PHASING_PCT, column=1, value="Phasing %")

    for i, period in enumerate(timeline.construction_months):
        col = col_letter(i)
        date_cell = ws[f"{col}{ROW_PHASING_DATE}"]
        date_cell.value = period.end
        date_cell.number_format = "mmm-yy"

        pct_cell = ws[f"{col}{ROW_PHASING_PCT}"]
        pct_cell.value = f"=Assumptions_Periodic_Capex!{col}{per_capex.ROW_ACTIVE}"
        pct_cell.font = Font(color=COLOR_LINK)
        pct_cell.number_format = "0.00%"

    ws.freeze_panes = ws.cell(row=ROW_PHASING_PCT + 1, column=FIRST_DATA_COL)

    return ws


def _resolved(ws: Worksheet, label_cell: str, label: str, value_cell: str,
              constant_row: int, number_format: str) -> None:
    ws[label_cell] = label
    cell = ws[value_cell]
    cell.value = f"=Assumptions_Constant!$B${constant_row}"
    cell.font = Font(color=COLOR_LINK)
    cell.number_format = number_format
