from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.workbook import Workbook

from inputs import ProjectInputs
from timeline import Timeline
from workbook_builder import (
    FIRST_DATA_COL,
    COLOR_INPUT,
    COLOR_FORMULA,
    TAB_COLOR_INPUT,
    col_letter,
)

N_SCENARIOS = 10

ROW_DATE_HEADER = 4
ROW_MONTH_INDEX = 5

ROW_BLOCK_TITLE = 7
ROW_FIRST_SCENARIO = 8                      # Scenario 1 .. 10 occupy rows 8-17
ROW_ACTIVE = ROW_FIRST_SCENARIO + N_SCENARIOS   # row 18

ROW_SUM_TITLE = 20
ROW_FIRST_SUM = 21                          # per-scenario phasing sums, rows 21-30

ROW_CHECK_HEADER = 33
ROW_CHECK_ACTIVE_SUMS = 34
ROW_CHECK_ALL_SUM = 35


def build_assumptions_periodic_capex(wb: Workbook, timeline: Timeline, inputs: ProjectInputs) -> Worksheet:
    ws = wb.create_sheet("Assumptions_Periodic_Capex")
    ws.sheet_properties.tabColor = TAB_COLOR_INPUT

    ws["A1"] = "Assumptions_Periodic_Capex — Monthly, Scenarios as Row Blocks"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = (
        "Time runs across columns, as everywhere else in the model; scenarios stack "
        "downward. The Active row resolves the selection on Cover."
    )
    ws["A2"].font = Font(italic=True, size=9)

    n_months = len(timeline.construction_months)
    first_col = col_letter(0)
    last_col = col_letter(n_months - 1)

    ws.cell(row=ROW_DATE_HEADER, column=1, value="Period End Date")
    ws.cell(row=ROW_MONTH_INDEX, column=1, value="Construction Month #")
    ws.cell(row=ROW_BLOCK_TITLE, column=1, value="Capex Phasing % (by scenario)").font = Font(bold=True)
    ws.cell(row=ROW_ACTIVE, column=1, value="ACTIVE — Capex Phasing %").font = Font(bold=True)

    for i, period in enumerate(timeline.construction_months):
        col = col_letter(i)
        date_cell = ws[f"{col}{ROW_DATE_HEADER}"]
        date_cell.value = period.end
        date_cell.number_format = "mmm-yy"
        ws[f"{col}{ROW_MONTH_INDEX}"] = i + 1

    for s in range(N_SCENARIOS):
        row = ROW_FIRST_SCENARIO + s
        ws.cell(row=row, column=1, value=f"Scenario {s + 1}")
        for i in range(n_months):
            col = col_letter(i)
            cell = ws[f"{col}{row}"]
            cell.value = inputs.capex.phasing_pct_by_month[i]
            cell.font = Font(color=COLOR_INPUT)
            cell.number_format = "0.00%"

    for i in range(n_months):
        col = col_letter(i)
        cell = ws[f"{col}{ROW_ACTIVE}"]
        cell.value = (
            f"=INDEX({col}${ROW_FIRST_SCENARIO}:{col}${ROW_FIRST_SCENARIO + N_SCENARIOS - 1},"
            f"Cover!$B$20)"
        )
        cell.font = Font(color=COLOR_FORMULA)
        cell.number_format = "0.00%"

    ws.cell(row=ROW_SUM_TITLE, column=1, value="Phasing sum by scenario (each must be 100%)").font = Font(bold=True)
    for s in range(N_SCENARIOS):
        row = ROW_FIRST_SUM + s
        ws.cell(row=row, column=1, value=f"Scenario {s + 1} sum")
        cell = ws.cell(row=row, column=2)
        cell.value = f"=SUM({first_col}{ROW_FIRST_SCENARIO + s}:{last_col}{ROW_FIRST_SCENARIO + s})"
        cell.font = Font(color=COLOR_FORMULA)
        cell.number_format = "0.0000%"

    ws.cell(row=ROW_CHECK_HEADER, column=1, value="Checks").font = Font(bold=True)
    ws.cell(row=ROW_CHECK_ACTIVE_SUMS, column=1, value="Check: Active phasing sums to 100%")
    ws.cell(row=ROW_CHECK_ALL_SUM, column=1, value="Check: All 10 scenario phasings sum to 100%")

    active_sum = ws.cell(row=ROW_CHECK_ACTIVE_SUMS, column=2)
    active_sum.value = f"=IF(ROUND(SUM({first_col}{ROW_ACTIVE}:{last_col}{ROW_ACTIVE})-1,6)=0,1,0)"
    active_sum.font = Font(color=COLOR_FORMULA)

    last_sum_row = ROW_FIRST_SUM + N_SCENARIOS - 1
    all_sum = ws.cell(row=ROW_CHECK_ALL_SUM, column=2)
    all_sum.value = (
        f"=IF(SUMPRODUCT(--(ROUND($B${ROW_FIRST_SUM}:$B${last_sum_row}-1,6)<>0))=0,1,0)"
    )
    all_sum.font = Font(color=COLOR_FORMULA)

    _name(wb, "PerCapex_ActiveSumCheck", "Assumptions_Periodic_Capex", f"B{ROW_CHECK_ACTIVE_SUMS}")
    _name(wb, "PerCapex_AllSumCheck", "Assumptions_Periodic_Capex", f"B{ROW_CHECK_ALL_SUM}")

    ws.column_dimensions["A"].width = 40
    ws.freeze_panes = ws.cell(row=ROW_BLOCK_TITLE + 1, column=FIRST_DATA_COL)

    return ws


def _name(wb: Workbook, name: str, sheet: str, cell: str) -> None:
    import re

    from openpyxl.workbook.defined_name import DefinedName

    m = re.match(r"([A-Za-z]+)(\d+)", cell)
    wb.defined_names.add(DefinedName(name, attr_text=f"'{sheet}'!${m.group(1)}${m.group(2)}"))
