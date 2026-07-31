from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.workbook import Workbook

import calc_financing_ops as fin_ops
import cover_refs as refs
import calc_revenue_opex as rev_opex
import calc_tax as tax
from inputs import ProjectInputs
from timeline import Timeline
from workbook_builder import (
    FIRST_DATA_COL,
    COLOR_FORMULA,
    COLOR_LINK,
    TAB_COLOR_CALC,
    col_letter,
)

ROW_DATE_HEADER = 2
ROW_QUARTER_INDEX = 3

ROW_CFADS = 5           # EBITDA - Tax, before debt service and before reserve movements
ROW_DEBT_SERVICE = 6    # linked from Calc_Financing_Ops
ROW_CFADS_POST_DS = 7

ROW_DSRA_CASH_COST = 9  # linked from Calc_Financing_Ops: funding, or LC fee when LC-backed

ROW_MAINT_CAPEX = 11    # linked from Calc_Revenue_Opex
ROW_MRA_TARGET = 12
ROW_MRA_BALANCE = 13
ROW_MRA_FUNDING = 14

ROW_CAFD = 16           # cash available for distribution, before the buffer and lock-up

# Distributions used to be 100% of CAFD including its negatives, which silently turned
# every shortfall into an equity call. The buffer absorbs the timing (the MRA traps a
# whole overhaul four quarters before it lands), the lock-up blocks distributions on weak
# DSCR, and anything still short shows up on its own line rather than as a negative
# dividend.
ROW_DSCR = 18
ROW_LOCKUP_FLAG = 19
ROW_BUFFER_TARGET = 20
ROW_BUFFER_OPENING = 21
ROW_DISTRIBUTION = 22
ROW_EQUITY_INJECTION = 23
ROW_BUFFER_CLOSING = 24

ROW_FCFE = 26           # distributions less equity injections
ROW_FCFF = 27           # unlevered: CFADS less maintenance capex, built once and reused

ROW_CHECK_HEADER = 30
ROW_CHECK_FCFE_NOT_BELOW_ZERO_COUNT = 31  # informational, not a hard fail
ROW_CHECK_MRA_WINDS_DOWN = 32
ROW_CHECK_MRA_FUNDS_CAPEX = 33
ROW_CHECK_BUFFER_NON_NEGATIVE = 34
ROW_CHECK_CASH_RECONCILES = 35
ROW_CHECK_EQUITY_INJECTIONS = 36          # informational: how many quarters needed one


def build_calc_cfads(wb: Workbook, timeline: Timeline, inputs: ProjectInputs) -> Worksheet:
    ws = wb.create_sheet("Calc_CFADS")
    ws.sheet_properties.tabColor = TAB_COLOR_CALC

    ws["A1"] = "Calc_CFADS — Quarterly Cash Waterfall (DSRA cash cost, MRA, maintenance capex)"
    ws["A1"].font = Font(bold=True, size=12)

    n_quarters = len(timeline.operations_quarters)
    mra_quarters = inputs.reserves.mra_lookforward_quarters

    _label(ws, ROW_DATE_HEADER, "Period End Date")
    _label(ws, ROW_QUARTER_INDEX, "Operating Quarter #")
    _label(ws, ROW_CFADS, "CFADS ($) = EBITDA - Tax")
    _label(ws, ROW_DEBT_SERVICE, "Debt Service ($) — linked from Calc_Financing_Ops")
    _label(ws, ROW_CFADS_POST_DS, "CFADS after Debt Service ($)")
    _label(ws, ROW_DSRA_CASH_COST, "DSRA Funding/(Release) or LC Fee ($) — linked from Calc_Financing_Ops")
    _label(ws, ROW_MAINT_CAPEX, "Maintenance Capex ($) — linked from Calc_Revenue_Opex")
    _label(ws, ROW_MRA_TARGET, f"MRA Target ($) = next {mra_quarters} quarters' maintenance capex")
    _label(ws, ROW_MRA_BALANCE, "MRA Balance ($)")
    _label(ws, ROW_MRA_FUNDING, "MRA Funding/(Release) ($) — releases as the spend it pre-funded lands")
    _label(ws, ROW_CAFD, "Cash Available for Distribution ($) — after debt service and reserves")
    _label(ws, ROW_DSCR, "DSCR (achieved) — linked from Calc_Financing_Ops")
    _label(ws, ROW_LOCKUP_FLAG, "Distributions Locked Up? (1 = blocked by DSCR trigger)")
    _label(ws, ROW_BUFFER_TARGET, "Target Cash Buffer ($) = target quarters x quarterly opex")
    _label(ws, ROW_BUFFER_OPENING, "Cash Buffer, Opening ($)")
    _label(ws, ROW_DISTRIBUTION, "Distribution to Equity ($)")
    _label(ws, ROW_EQUITY_INJECTION, "Equity Injection Required ($) — shortfall the buffer cannot cover")
    _label(ws, ROW_BUFFER_CLOSING, "Cash Buffer, Closing ($)")
    _label(ws, ROW_FCFE, "FCFE ($) = Distributions less Equity Injections")
    _label(ws, ROW_FCFF, "FCFF ($) = CFADS - Maintenance Capex (unlevered)")

    ws.cell(row=ROW_CHECK_HEADER, column=1, value="Checks").font = Font(bold=True)
    _label(ws, ROW_CHECK_FCFE_NOT_BELOW_ZERO_COUNT, "Informational: # of quarters with negative FCFE")
    _label(ws, ROW_CHECK_MRA_WINDS_DOWN, "Check: MRA fully released by end of life (net funding = 0)")
    _label(ws, ROW_CHECK_MRA_FUNDS_CAPEX, "Check: MRA pre-funds each quarter's spend (prior balance >= capex)")
    _label(ws, ROW_CHECK_BUFFER_NON_NEGATIVE, "Check: Cash buffer never negative")
    _label(ws, ROW_CHECK_CASH_RECONCILES, "Check: Sum of CAFD = distributions - injections + closing buffer")
    _label(ws, ROW_CHECK_EQUITY_INJECTIONS, "Informational: # of quarters needing an equity injection")

    for i, period in enumerate(timeline.operations_quarters):
        col = col_letter(i)
        prev = col_letter(i - 1) if i > 0 else None

        ws[f"{col}{ROW_DATE_HEADER}"] = period.end
        ws[f"{col}{ROW_DATE_HEADER}"].number_format = "mmm-yy"
        ws[f"{col}{ROW_QUARTER_INDEX}"] = i + 1

        # CFADS stays pre-maintenance-capex: it is the stream the sculpting basis divides
        # by Target DSCR, and maintenance spend is funded through the MRA line below.
        cfads_cell = ws[f"{col}{ROW_CFADS}"]
        cfads_cell.value = f"=Calc_Tax!{col}{tax.ROW_EBITDA}-Calc_Tax!{col}{tax.ROW_TAX}"
        cfads_cell.font = Font(color=COLOR_LINK)
        cfads_cell.number_format = "#,##0"

        _link(ws, col, ROW_DEBT_SERVICE, f"Calc_Financing_Ops!{col}{fin_ops.ROW_DEBT_SERVICE}")
        _formula(ws, col, ROW_CFADS_POST_DS, f"={col}{ROW_CFADS}-{col}{ROW_DEBT_SERVICE}")

        _link(ws, col, ROW_DSRA_CASH_COST, f"Calc_Financing_Ops!{col}{fin_ops.ROW_DSRA_CASH_COST}")

        _link(ws, col, ROW_MAINT_CAPEX, f"Calc_Revenue_Opex!{col}{rev_opex.ROW_MAINT_TOTAL}")

        if i + 1 < n_quarters:
            window_start = col_letter(i + 1)
            window_end = col_letter(min(i + mra_quarters, n_quarters - 1))
            _formula(ws, col, ROW_MRA_TARGET,
                     f"=SUM({window_start}{ROW_MAINT_CAPEX}:{window_end}{ROW_MAINT_CAPEX})")
        else:
            _formula(ws, col, ROW_MRA_TARGET, "=0")

        _formula(ws, col, ROW_MRA_BALANCE, f"={col}{ROW_MRA_TARGET}")

        if prev is None:
            _formula(ws, col, ROW_MRA_FUNDING, f"={col}{ROW_MRA_BALANCE}")
        else:
            _formula(ws, col, ROW_MRA_FUNDING, f"={col}{ROW_MRA_BALANCE}-{prev}{ROW_MRA_BALANCE}")

        # Maintenance capex is a separate line from the MRA movement: the reserve releases
        # cash in the quarter the spend lands, and netting the two would hide both.
        _formula(ws, col, ROW_CAFD,
                 f"={col}{ROW_CFADS_POST_DS}-{col}{ROW_DSRA_CASH_COST}"
                 f"-{col}{ROW_MRA_FUNDING}-{col}{ROW_MAINT_CAPEX}")

        dscr_cell = ws[f"{col}{ROW_DSCR}"]
        dscr_cell.value = f"=Calc_Financing_Ops!{col}{fin_ops.ROW_DSCR}"
        dscr_cell.font = Font(color=COLOR_LINK)
        dscr_cell.number_format = "0.00x"

        # Post-tenor the DSCR row is blank, which must read as "not locked" rather than
        # as a zero that traps cash for the rest of the project's life.
        _formula(ws, col, ROW_LOCKUP_FLAG,
                 f"=IF(N({col}{ROW_DSCR})=0,0,IF({col}{ROW_DSCR}<Cover!{refs.ABS_LOCKUP_DSCR},1,0))")

        # There is no going concern past the horizon, so the final quarter holds nothing
        # back. Without this the buffer is stranded at end of life — value that was never
        # returned to equity, quietly depressing EIRR.
        if i == n_quarters - 1:
            _formula(ws, col, ROW_BUFFER_TARGET, "=0")
        else:
            _formula(ws, col, ROW_BUFFER_TARGET,
                     f"=Cover!{refs.ABS_CASH_BUFFER_TARGET}*Calc_Revenue_Opex!{col}{rev_opex.ROW_OPEX}")

        if prev is None:
            _formula(ws, col, ROW_BUFFER_OPENING, "=0")
        else:
            _formula(ws, col, ROW_BUFFER_OPENING, f"={prev}{ROW_BUFFER_CLOSING}")

        available = f"({col}{ROW_BUFFER_OPENING}+{col}{ROW_CAFD})"
        # Distribute only what is above the target buffer, and nothing at all while locked
        # up. Retaining to the target is what gives later quarters something to draw on.
        _formula(ws, col, ROW_DISTRIBUTION,
                 f"=IF({col}{ROW_LOCKUP_FLAG}=1,0,"
                 f"MAX(0,{available}-{col}{ROW_BUFFER_TARGET}))")
        # Whatever the buffer still cannot cover is a genuine call on shareholders. Naming
        # it is the whole point — before this it hid inside a negative dividend.
        _formula(ws, col, ROW_EQUITY_INJECTION, f"=MAX(0,-{available})")
        _formula(ws, col, ROW_BUFFER_CLOSING,
                 f"={available}-{col}{ROW_DISTRIBUTION}+{col}{ROW_EQUITY_INJECTION}")

        _formula(ws, col, ROW_FCFE,
                 f"={col}{ROW_DISTRIBUTION}-{col}{ROW_EQUITY_INJECTION}")
        _formula(ws, col, ROW_FCFF, f"={col}{ROW_CFADS}-{col}{ROW_MAINT_CAPEX}")

    last_col = col_letter(n_quarters - 1)
    first_col = col_letter(0)

    check_cell = ws[f"{last_col}{ROW_CHECK_FCFE_NOT_BELOW_ZERO_COUNT}"]
    check_cell.value = f'=COUNTIF({first_col}{ROW_FCFE}:{last_col}{ROW_FCFE},"<0")'
    check_cell.font = Font(color=COLOR_FORMULA)

    # Every dollar trapped must come back out: the fundings and releases net to zero over
    # the life, which is what makes the reserve a timing device rather than a leak.
    winds_down = ws[f"{last_col}{ROW_CHECK_MRA_WINDS_DOWN}"]
    winds_down.value = (
        f"=IF(ROUND(SUM({first_col}{ROW_MRA_FUNDING}:{last_col}{ROW_MRA_FUNDING}),2)=0,1,0)"
    )
    winds_down.font = Font(color=COLOR_FORMULA)

    # The reserve must be standing before the spend arrives, not after: comparing each
    # quarter's capex against the *previous* quarter's balance is what catches a lookforward
    # window wired backwards, which nothing else here would notice.
    second_col = col_letter(1)
    penultimate_col = col_letter(n_quarters - 2)
    funds_capex = ws[f"{last_col}{ROW_CHECK_MRA_FUNDS_CAPEX}"]
    funds_capex.value = (
        f"=IF(SUMPRODUCT(--({first_col}{ROW_MRA_BALANCE}:{penultimate_col}{ROW_MRA_BALANCE}"
        f"<{second_col}{ROW_MAINT_CAPEX}:{last_col}{ROW_MAINT_CAPEX}-0.01))=0,1,0)"
    )
    funds_capex.font = Font(color=COLOR_FORMULA)

    buffer_ok = ws[f"{last_col}{ROW_CHECK_BUFFER_NON_NEGATIVE}"]
    buffer_ok.value = (
        f"=IF(MIN({first_col}{ROW_BUFFER_CLOSING}:{last_col}{ROW_BUFFER_CLOSING})>=-0.01,1,0)"
    )
    buffer_ok.font = Font(color=COLOR_FORMULA)

    # Every dollar of CAFD is either distributed, offset by an injection, or still sitting
    # in the buffer. If those three do not tie, cash is being created or destroyed.
    reconciles = ws[f"{last_col}{ROW_CHECK_CASH_RECONCILES}"]
    reconciles.value = (
        f"=IF(ROUND(SUM({first_col}{ROW_CAFD}:{last_col}{ROW_CAFD})"
        f"-SUM({first_col}{ROW_DISTRIBUTION}:{last_col}{ROW_DISTRIBUTION})"
        f"+SUM({first_col}{ROW_EQUITY_INJECTION}:{last_col}{ROW_EQUITY_INJECTION})"
        f"-{last_col}{ROW_BUFFER_CLOSING},2)=0,1,0)"
    )
    reconciles.font = Font(color=COLOR_FORMULA)

    injections = ws[f"{last_col}{ROW_CHECK_EQUITY_INJECTIONS}"]
    injections.value = (
        f'=COUNTIF({first_col}{ROW_EQUITY_INJECTION}:{last_col}{ROW_EQUITY_INJECTION},">0.01")'
    )
    injections.font = Font(color=COLOR_FORMULA)

    _add_named_range(wb, "CFADS_LastCol", "Calc_CFADS", f"{last_col}1")
    _add_named_range(wb, "CFADS_NegFCFECount", "Calc_CFADS", f"{last_col}{ROW_CHECK_FCFE_NOT_BELOW_ZERO_COUNT}")
    _add_named_range(wb, "CFADS_MRAWindsDownCheck", "Calc_CFADS", f"{last_col}{ROW_CHECK_MRA_WINDS_DOWN}")
    _add_named_range(wb, "CFADS_MRAFundsCapexCheck", "Calc_CFADS", f"{last_col}{ROW_CHECK_MRA_FUNDS_CAPEX}")
    _add_named_range(wb, "CFADS_BufferNonNegCheck", "Calc_CFADS", f"{last_col}{ROW_CHECK_BUFFER_NON_NEGATIVE}")
    _add_named_range(wb, "CFADS_CashReconcilesCheck", "Calc_CFADS", f"{last_col}{ROW_CHECK_CASH_RECONCILES}")
    _add_named_range(wb, "CFADS_EquityInjectionCount", "Calc_CFADS", f"{last_col}{ROW_CHECK_EQUITY_INJECTIONS}")

    ws.freeze_panes = ws.cell(row=ROW_FCFF + 1, column=FIRST_DATA_COL)
    ws.column_dimensions["A"].width = 58

    return ws


def _label(ws: Worksheet, row: int, label: str) -> None:
    ws.cell(row=row, column=1, value=label)


def _link(ws: Worksheet, col: str, row: int, formula: str) -> None:
    cell = ws[f"{col}{row}"]
    cell.value = f"={formula}"
    cell.font = Font(color=COLOR_LINK)
    cell.number_format = "#,##0"


def _formula(ws: Worksheet, col: str, row: int, formula: str) -> None:
    cell = ws[f"{col}{row}"]
    cell.value = formula
    cell.font = Font(color=COLOR_FORMULA)
    cell.number_format = "#,##0"


def _add_named_range(wb: Workbook, name: str, sheet: str, cell: str) -> None:
    import re

    from openpyxl.workbook.defined_name import DefinedName

    match = re.match(r"([A-Za-z]+)(\d+)", cell)
    col, row = match.group(1), match.group(2)
    wb.defined_names.add(DefinedName(name, attr_text=f"'{sheet}'!${col}${row}"))
