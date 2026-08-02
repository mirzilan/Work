from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.workbook import Workbook

import assumptions_model as model
import assumptions_periodic_ops as per_ops
from inputs import ProjectInputs
from timeline import Timeline
from workbook_builder import (
    FIRST_DATA_COL,
    COLOR_FORMULA,
    COLOR_LINK,
    TAB_COLOR_CALC,
    col_letter,
)

ROW_DATE_HEADER = 2
ROW_QUARTER_INDEX = 3

ROW_REV_VOLUME_INDEX = 5
ROW_REV_ESC_INDEX = 6
ROW_OPEX_VOLUME_INDEX = 7
ROW_OPEX_ESC_INDEX = 8

ROW_REVENUE = 10
ROW_OPEX = 11
ROW_EBITDA = 12

# Maintenance capex sits here rather than on Calc_CFADS because it is revenue-driven and
# must be readable by Calc_Tax (depreciation vintages). Sourcing it from the waterfall
# sheet would put Tax and CFADS in a cycle that no staged cell breaks.
ROW_MAINT_ROUTINE = 14
ROW_MAINT_LUMPY = 15
ROW_MAINT_TOTAL = 16

ROW_CHECK_HEADER = 19
ROW_CHECK_YEAR1_REVENUE = 20
ROW_CHECK_EBITDA_POSITIVE = 21
ROW_CHECK_MAINT_NONNEG = 22

ANNUAL_REVENUE_CELL = "B4"
OPEX_PCT_CELL = "B9"
ROUTINE_MAINT_PCT_CELL = "B13"
ABS_ROUTINE_MAINT_PCT = "$B$13"


def build_calc_revenue_opex(wb: Workbook, timeline: Timeline, inputs: ProjectInputs) -> Worksheet:
    ws = wb.create_sheet("Calc_Revenue_Opex")
    ws.sheet_properties.tabColor = TAB_COLOR_CALC

    ws["A1"] = "Calc_Revenue_Opex — Quarterly Revenue & Opex (flat dummy, Stage 1a)"
    ws["A1"].font = Font(bold=True, size=12)

    ws["A4"] = "Annual Revenue Input ($) — linked from Assumptions_Model, flat dummy placeholder"
    ws[ANNUAL_REVENUE_CELL] = f"=Assumptions_Model!$B${model.ROW_ANNUAL_REVENUE}"
    ws[ANNUAL_REVENUE_CELL].font = Font(color=COLOR_LINK)
    ws[ANNUAL_REVENUE_CELL].number_format = "#,##0"

    ws["A9"] = "Opex % of Revenue — linked from Assumptions_Model"
    ws[OPEX_PCT_CELL] = f"=Assumptions_Model!$B${model.ROW_OPEX_PCT}"
    ws[OPEX_PCT_CELL].font = Font(color=COLOR_LINK)
    ws[OPEX_PCT_CELL].number_format = "0.00%"

    ws["A13"] = "Routine Maint Capex (% of Revenue) — linked from Assumptions_Model"
    ws[ROUTINE_MAINT_PCT_CELL] = f"=Assumptions_Model!$B${model.ROW_ROUTINE_MAINT_PCT}"
    ws[ROUTINE_MAINT_PCT_CELL].font = Font(color=COLOR_LINK)
    ws[ROUTINE_MAINT_PCT_CELL].number_format = "0.00%"

    _label(ws, ROW_DATE_HEADER, "Period End Date")
    _label(ws, ROW_QUARTER_INDEX, "Operating Quarter #")
    _label(ws, ROW_REV_VOLUME_INDEX, "Revenue Volume Index — linked from Assumptions_Periodic_Ops")
    _label(ws, ROW_REV_ESC_INDEX, "Revenue Escalation Index — linked from Assumptions_Periodic_Ops")
    _label(ws, ROW_OPEX_VOLUME_INDEX, "Opex Volume Index — linked from Assumptions_Periodic_Ops")
    _label(ws, ROW_OPEX_ESC_INDEX, "Opex Escalation Index — linked from Assumptions_Periodic_Ops")
    _label(ws, ROW_REVENUE, "Revenue ($) = base x volume index x escalation index")
    _label(ws, ROW_OPEX, "Opex ($) = base x volume index x escalation index")
    _label(ws, ROW_EBITDA, "EBITDA ($)")
    _label(ws, ROW_MAINT_ROUTINE, "Routine Maintenance Capex ($) = Revenue x routine %")
    _label(ws, ROW_MAINT_LUMPY, "Lumpy Maintenance Capex ($) — linked from Assumptions_Periodic_Ops")
    _label(ws, ROW_MAINT_TOTAL, "Total Maintenance Capex ($)")

    ws.cell(row=ROW_CHECK_HEADER, column=1, value="Checks").font = Font(bold=True)
    _label(ws, ROW_CHECK_YEAR1_REVENUE, "Check: Year 1 revenue = Annual Revenue Input (holds when Yr1 indices = 1.00)")
    _label(ws, ROW_CHECK_EBITDA_POSITIVE, "Informational: # quarters with negative EBITDA")
    _label(ws, ROW_CHECK_MAINT_NONNEG, "Check: Total maintenance capex >= 0 in every quarter")

    n_quarters = len(timeline.operations_quarters)

    for i, period in enumerate(timeline.operations_quarters):
        col = col_letter(i)

        ws[f"{col}{ROW_DATE_HEADER}"] = period.end
        ws[f"{col}{ROW_DATE_HEADER}"].number_format = "mmm-yy"
        ws[f"{col}{ROW_QUARTER_INDEX}"] = i + 1

        for row, source_row in (
            (ROW_REV_VOLUME_INDEX, per_ops.ROW_REV_ACTIVE),
            (ROW_REV_ESC_INDEX, per_ops.ROW_REV_ESC_ACTIVE),
            (ROW_OPEX_VOLUME_INDEX, per_ops.ROW_OPEX_ACTIVE),
            (ROW_OPEX_ESC_INDEX, per_ops.ROW_OPEX_ESC_ACTIVE),
        ):
            idx_cell = ws[f"{col}{row}"]
            idx_cell.value = f"=Assumptions_Periodic_Ops!{col}{source_row}"
            idx_cell.font = Font(color=COLOR_LINK)
            idx_cell.number_format = "0.0000"

        rev_cell = ws[f"{col}{ROW_REVENUE}"]
        rev_cell.value = f"=$B$4/4*{col}{ROW_REV_VOLUME_INDEX}*{col}{ROW_REV_ESC_INDEX}"
        rev_cell.font = Font(color=COLOR_FORMULA)
        rev_cell.number_format = "#,##0"

        # Opex escalates off its own base rather than off already-escalated revenue —
        # otherwise revenue escalation would be double-counted in the cost line.
        opex_cell = ws[f"{col}{ROW_OPEX}"]
        opex_cell.value = f"=$B$4/4*$B$9*{col}{ROW_OPEX_VOLUME_INDEX}*{col}{ROW_OPEX_ESC_INDEX}"
        opex_cell.font = Font(color=COLOR_FORMULA)
        opex_cell.number_format = "#,##0"

        ebitda_cell = ws[f"{col}{ROW_EBITDA}"]
        ebitda_cell.value = f"={col}{ROW_REVENUE}-{col}{ROW_OPEX}"
        ebitda_cell.font = Font(color=COLOR_FORMULA)
        ebitda_cell.number_format = "#,##0"

        routine_cell = ws[f"{col}{ROW_MAINT_ROUTINE}"]
        routine_cell.value = f"={col}{ROW_REVENUE}*{ABS_ROUTINE_MAINT_PCT}"
        routine_cell.font = Font(color=COLOR_FORMULA)
        routine_cell.number_format = "#,##0"

        lumpy_cell = ws[f"{col}{ROW_MAINT_LUMPY}"]
        lumpy_cell.value = f"=Assumptions_Periodic_Ops!{col}{per_ops.ROW_MAINT_ACTIVE}"
        lumpy_cell.font = Font(color=COLOR_LINK)
        lumpy_cell.number_format = "#,##0"

        maint_total_cell = ws[f"{col}{ROW_MAINT_TOTAL}"]
        maint_total_cell.value = f"={col}{ROW_MAINT_ROUTINE}+{col}{ROW_MAINT_LUMPY}"
        maint_total_cell.font = Font(color=COLOR_FORMULA)
        maint_total_cell.number_format = "#,##0"

    # Check: first 4 quarters' revenue sums to the annual input
    q1, q2, q3, q4 = (col_letter(i) for i in range(4))
    check_cell = ws[f"{q4}{ROW_CHECK_YEAR1_REVENUE}"]
    check_cell.value = (
        f"=IF(ROUND(SUM({q1}{ROW_REVENUE}:{q4}{ROW_REVENUE})-$B$4,2)=0,1,0)"
    )
    check_cell.font = Font(color=COLOR_FORMULA)

    last_col = col_letter(n_quarters - 1)
    first_col = col_letter(0)

    ebitda_check = ws[f"{last_col}{ROW_CHECK_EBITDA_POSITIVE}"]
    ebitda_check.value = f"=COUNTIF({first_col}{ROW_EBITDA}:{last_col}{ROW_EBITDA},\"<0\")"
    ebitda_check.font = Font(color=COLOR_FORMULA)

    maint_check = ws[f"{last_col}{ROW_CHECK_MAINT_NONNEG}"]
    maint_check.value = f"=IF(MIN({first_col}{ROW_MAINT_TOTAL}:{last_col}{ROW_MAINT_TOTAL})>=0,1,0)"
    maint_check.font = Font(color=COLOR_FORMULA)

    _add_named_range(wb, "RevOpex_LastCol", "Calc_Revenue_Opex", f"{last_col}1")
    _add_named_range(wb, "RevOpex_Year1Check", "Calc_Revenue_Opex", f"{q4}{ROW_CHECK_YEAR1_REVENUE}")
    _add_named_range(wb, "RevOpex_MaintNonNegCheck", "Calc_Revenue_Opex", f"{last_col}{ROW_CHECK_MAINT_NONNEG}")

    ws.freeze_panes = ws.cell(row=ROW_MAINT_TOTAL + 1, column=FIRST_DATA_COL)

    return ws


def _label(ws: Worksheet, row: int, label: str) -> None:
    ws.cell(row=row, column=1, value=label)


def _add_named_range(wb: Workbook, name: str, sheet: str, cell: str) -> None:
    import re

    from openpyxl.workbook.defined_name import DefinedName

    match = re.match(r"([A-Za-z]+)(\d+)", cell)
    col, row = match.group(1), match.group(2)
    wb.defined_names.add(DefinedName(name, attr_text=f"'{sheet}'!${col}${row}"))
