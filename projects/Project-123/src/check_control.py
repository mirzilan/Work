from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.workbook import Workbook

import assumptions_constant as const
import assumptions_periodic_capex as per_capex
import assumptions_periodic_ops as per_ops
import calc_capex as capex
import calc_cfads as cfads
import calc_financing_cons as fin_cons
import calc_financing_ops as fin_ops
import calc_revenue_opex as rev_opex
import calc_tax as tax
import cover
import fs_quarterly as fsq
import valuation_selldown as selldown
from timeline import Timeline
from workbook_builder import TAB_COLOR_CHECK, COLOR_FORMULA, col_letter

ROW_HEADER = 1
ROW_MASTER_FLAG = 3
ROW_TABLE_HEADER = 5
ROW_FIRST_CHECK = 6

# "flag"  — cell is 1/0, 1 = pass -> OK/FAIL
# "count" — cell is a failure count, 0 = pass -> OK/FAIL
# "info"  — cell is a count that's informational only -> OK/REVIEW, never FAIL
FLAG, COUNT, INFO = "flag", "count", "info"


def _checks(timeline: Timeline) -> list[tuple[str, str, str, str]]:
    """(sheet, description, cell, check_type). Every cell reference is derived from the
    source module's own row constants — a row moving on a calc sheet must not be able to
    leave this aggregator silently pointing at a blank cell."""
    cons = col_letter(len(timeline.construction_months) - 1)
    ops = col_letter(len(timeline.operations_quarters) - 1)
    q4 = col_letter(3)

    return [
        ("Cover", "Solve is current (assumptions unchanged since last solve)",
         f"D{cover.ROW_FRESHNESS_FLAG}", FLAG),

        ("Calc_Capex", "Cum Debt + Cum Equity = Cum Capex + Cum IDC + Initial DSRA",
         f"{cons}{capex.ROW_CHECK_FUNDING_TIES}", FLAG),
        ("Calc_Capex", "Cumulative Capex = Total Capex Input",
         f"{cons}{capex.ROW_CHECK_TOTAL_MATCHES_INPUT}", FLAG),

        ("Calc_Financing_Cons", "Closing Balance = Cumulative Debt Draws",
         f"{cons}{fin_cons.ROW_CHECK_CLOSING_MATCHES_DRAWS}", FLAG),
        ("Calc_Financing_Cons", "IDC Converged (Loop 1)",
         f"{cons}{fin_cons.ROW_CHECK_IDC_CONVERGED}", FLAG),
        ("Calc_Financing_Cons", "Cum Debt + Cum Equity = Cum Funding Requirement",
         f"{cons}{fin_cons.ROW_CHECK_SOURCES_TIE_USES}", FLAG),
        ("Calc_Financing_Cons", "Cumulative Debt Draw <= Debt Facility",
         f"{cons}{fin_cons.ROW_CHECK_WITHIN_FACILITY}", FLAG),
        ("Calc_Financing_Cons", "Debt draws honour the solved facility",
         f"{cons}{fin_cons.ROW_CHECK_FACILITY_FULLY_DRAWN}", FLAG),

        ("Assumptions_Constant", "Active column resolves (Total Capex > 0)",
         f"B{const.ROW_CHECK_ACTIVE_RESOLVES}", FLAG),
        ("Assumptions_Constant", "All scenario cells populated",
         f"B{const.ROW_CHECK_ALL_POPULATED}", FLAG),

        ("Assumptions_Periodic_Capex", "Active phasing sums to 100%",
         f"B{per_capex.ROW_CHECK_ACTIVE_SUMS}", FLAG),
        ("Assumptions_Periodic_Capex", "All 10 scenario phasings sum to 100%",
         f"B{per_capex.ROW_CHECK_ALL_SUM}", FLAG),

        ("Assumptions_Periodic_Ops", "Escalation indices start at 1.00",
         f"B{per_ops.ROW_CHECK_ESC_BASE}", FLAG),
        ("Assumptions_Periodic_Ops", "Active volume indices > 0 every quarter",
         f"B{per_ops.ROW_CHECK_VOLUME_POSITIVE}", FLAG),
        ("Assumptions_Periodic_Ops", "Active lumpy maintenance capex >= 0 every quarter",
         f"B{per_ops.ROW_CHECK_MAINT_NONNEG}", FLAG),

        ("Calc_Revenue_Opex", "Year 1 Revenue Sum = Annual Revenue Input",
         f"{q4}{rev_opex.ROW_CHECK_YEAR1_REVENUE}", FLAG),
        ("Calc_Revenue_Opex", "Quarters with negative EBITDA",
         f"{ops}{rev_opex.ROW_CHECK_EBITDA_POSITIVE}", INFO),
        ("Calc_Revenue_Opex", "Total maintenance capex >= 0 every quarter",
         f"{ops}{rev_opex.ROW_CHECK_MAINT_NONNEG}", FLAG),

        ("Calc_Tax", "Accumulated Base Depreciation <= Total Project Cost",
         f"{ops}{tax.ROW_CHECK_ACCUM_DEPR}", FLAG),
        ("Calc_Tax", "Accumulated Maint Depreciation <= Cumulative Maint Capex",
         f"{ops}{tax.ROW_CHECK_ACCUM_MAINT_DEPR}", FLAG),
        ("Calc_Tax", "Tax losses c/f never negative",
         f"{ops}{tax.ROW_CHECK_TLCF_NON_NEGATIVE}", FLAG),
        ("Calc_Tax", "Losses utilised <= losses arising over life",
         f"{ops}{tax.ROW_CHECK_TLCF_RECONCILES}", FLAG),

        ("Calc_Financing_Ops", "Closing Balance = 0 at Debt Tenor End",
         f"{ops}{fin_ops.ROW_CHECK_FULLY_AMORTIZED}", FLAG),
        ("Calc_Financing_Ops", "Debt Sizing Converged (Loop 2)",
         f"{ops}{fin_ops.ROW_CHECK_SCULPT_CONVERGED}", FLAG),
        ("Calc_Financing_Ops", "Min DSCR over tenor >= Target DSCR",
         f"{ops}{fin_ops.ROW_CHECK_MIN_DSCR}", FLAG),
        ("Calc_Financing_Ops", "Quarters where CFADS/DSCR < interest",
         f"{ops}{fin_ops.ROW_CHECK_PRINCIPAL_FLOORED}", INFO),
        ("Calc_Financing_Ops", "Debt sizing capped by Max Gearing",
         f"{ops}{fin_ops.ROW_CHECK_CAPPED_BY_MAX_GEARING}", INFO),
        ("Calc_Financing_Ops", "PLCR >= LLCR in every quarter",
         f"{ops}{fin_ops.ROW_CHECK_PLCR_GE_LLCR}", FLAG),
        ("Calc_Financing_Ops", "LLCR >= 1.00 wherever debt is outstanding",
         f"{ops}{fin_ops.ROW_CHECK_MIN_LLCR}", FLAG),

        ("Calc_CFADS", "Negative FCFE Quarter Count",
         f"{ops}{cfads.ROW_CHECK_FCFE_NOT_BELOW_ZERO_COUNT}", INFO),
        ("Calc_CFADS", "MRA fully released by end of life (net funding = 0)",
         f"{ops}{cfads.ROW_CHECK_MRA_WINDS_DOWN}", FLAG),
        ("Calc_CFADS", "MRA pre-funds each quarter's maintenance spend",
         f"{ops}{cfads.ROW_CHECK_MRA_FUNDS_CAPEX}", FLAG),

        ("Calc_CFADS", "Cash buffer never negative",
         f"{ops}{cfads.ROW_CHECK_BUFFER_NON_NEGATIVE}", FLAG),
        ("Calc_CFADS", "CAFD = distributions - injections + closing buffer",
         f"{ops}{cfads.ROW_CHECK_CASH_RECONCILES}", FLAG),
        ("Calc_CFADS", "Quarters needing an equity injection",
         f"{ops}{cfads.ROW_CHECK_EQUITY_INJECTIONS}", INFO),

        ("FS_Quarterly", "BS Balance Failures (Assets != Liab+Equity)",
         f"{ops}{fsq.ROW_CHECK_BS_BALANCES_COUNT}", COUNT),
        ("FS_Quarterly", "Closing cash ties to the Calc_CFADS buffer",
         f"{ops}{fsq.ROW_CHECK_CASH_TIES_BUFFER}", FLAG),

        ("Valuation_SellDown", "Sale price is 0 at the final exit year",
         f"B{selldown.ROW_CHECK_FINAL_EXIT_ZERO_PRICE}", FLAG),
    ]


def build_check_control(wb: Workbook, timeline: Timeline) -> Worksheet:
    ws = wb.create_sheet("Check_Control")
    ws.sheet_properties.tabColor = TAB_COLOR_CHECK

    checks = _checks(timeline)

    ws["A1"] = "Check_Control — Master Audit Aggregator"
    ws["A1"].font = Font(bold=True, size=12)

    ws.cell(row=ROW_MASTER_FLAG, column=1, value="MASTER STATUS")
    ws.cell(row=ROW_MASTER_FLAG, column=1).font = Font(bold=True)

    for col, header in ((1, "Sheet"), (2, "Check Description"), (3, "Source Cell"),
                        (4, "Value"), (5, "Status")):
        cell = ws.cell(row=ROW_TABLE_HEADER, column=col, value=header)
        cell.font = Font(bold=True)

    row = ROW_FIRST_CHECK
    for sheet, desc, cell_ref, check_type in checks:
        ws.cell(row=row, column=1, value=sheet)
        ws.cell(row=row, column=2, value=desc)
        ws.cell(row=row, column=3, value=f"'{sheet}'!{cell_ref}")

        value_cell = ws.cell(row=row, column=4)
        value_cell.value = f"='{sheet}'!{cell_ref}"
        value_cell.font = Font(color=COLOR_FORMULA)

        status_cell = ws.cell(row=row, column=5)
        if check_type == FLAG:
            status_cell.value = f'=IF(D{row}=1,"OK","FAIL")'
        elif check_type == COUNT:
            status_cell.value = f'=IF(D{row}=0,"OK","FAIL")'
        else:
            status_cell.value = f'=IF(D{row}=0,"OK","REVIEW")'
        status_cell.font = Font(color=COLOR_FORMULA)

        row += 1

    last_check_row = row - 1

    master_cell = ws.cell(row=ROW_MASTER_FLAG, column=2)
    master_cell.value = (
        f'=IF(COUNTIF(E{ROW_FIRST_CHECK}:E{last_check_row},"FAIL")=0,"MODEL OK","ERRORS FOUND")'
    )
    master_cell.font = Font(bold=True, size=14, color=COLOR_FORMULA)

    _add_named_range(wb, "CheckControl_MasterFlag", "Check_Control", "B3")

    ws.column_dimensions["B"].width = 50
    ws.column_dimensions["C"].width = 26

    return ws


def _add_named_range(wb: Workbook, name: str, sheet: str, cell: str) -> None:
    import re

    from openpyxl.workbook.defined_name import DefinedName

    match = re.match(r"([A-Za-z]+)(\d+)", cell)
    col, row = match.group(1), match.group(2)
    wb.defined_names.add(DefinedName(name, attr_text=f"'{sheet}'!${col}${row}"))
