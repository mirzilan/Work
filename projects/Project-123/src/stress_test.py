from openpyxl.styles import Font, PatternFill
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.workbook import Workbook

import assumptions_constant as const
from workbook_builder import COLOR_INPUT, COLOR_FORMULA, TAB_COLOR_INPUT

# One-at-a-time sensitivity: each row below shocks a single Assumptions_Constant driver on
# the ACTIVE scenario only, re-solves both loops, records the result, then restores the
# original value before moving to the next row — never cumulative, so each row isolates
# one variable's effect. The macro (mod_StressTest.bas) reads this table generically
# rather than hardcoding the variable list, the same "macros address the model through
# named ranges only" principle the rest of the VBA already follows.

ROW_HEADER = 1
ROW_SPEC_HEADER = 3
ROW_SPEC_TABLE_HEADER = 4
ROW_FIRST_SPEC = 5

# (label, Assumptions_Constant row, shock type, shock value)
# "relative" multiplies the base value by (1 + shock); "absolute" adds the shock directly
# (used for rates/ratios where a % move, not a %-of-%, is the meaningful stress).
SHOCK_SPECS = [
    ("Total Capex +10%", const.ROW_TOTAL_CAPEX, "relative", 0.10),
    ("Total Capex -10%", const.ROW_TOTAL_CAPEX, "relative", -0.10),
    ("Annual Revenue +10%", const.ROW_ANNUAL_REVENUE, "relative", 0.10),
    ("Annual Revenue -10%", const.ROW_ANNUAL_REVENUE, "relative", -0.10),
    ("Opex % of Revenue +10%", const.ROW_OPEX_PCT, "relative", 0.10),
    ("Opex % of Revenue -10%", const.ROW_OPEX_PCT, "relative", -0.10),
    ("Interest Rate +100bps", const.ROW_INTEREST_RATE, "absolute", 0.01),
    ("Interest Rate -100bps", const.ROW_INTEREST_RATE, "absolute", -0.01),
    ("Target DSCR +0.10x", const.ROW_TARGET_DSCR, "absolute", 0.10),
    ("Target DSCR -0.10x", const.ROW_TARGET_DSCR, "absolute", -0.10),
]

N_SHOCKS = len(SHOCK_SPECS)
ROW_LAST_SPEC = ROW_FIRST_SPEC + N_SHOCKS - 1

ROW_RESULTS_HEADER = ROW_LAST_SPEC + 2
ROW_RESULTS_TABLE_HEADER = ROW_RESULTS_HEADER + 1
ROW_FIRST_RESULT = ROW_RESULTS_TABLE_HEADER + 1
ROW_LAST_RESULT = ROW_FIRST_RESULT + N_SHOCKS - 1

ROW_META_HEADER = ROW_LAST_RESULT + 2
ROW_LAST_RUN = ROW_META_HEADER + 1
ROW_RUN_SCENARIO = ROW_META_HEADER + 2

# Results columns (1-indexed, matching the macro's row.Offset(0, n) writes)
COL_VARIABLE = 1
COL_SHOCK_DISPLAY = 2
COL_BASE_EIRR = 3
COL_STRESSED_EIRR = 4
COL_DELTA_EIRR_BPS = 5
COL_BASE_PIRR = 6
COL_STRESSED_PIRR = 7
COL_DELTA_PIRR_BPS = 8
COL_GEARING = 9
COL_MIN_DSCR = 10
COL_CONVERGED = 11

FILL_GREEN = PatternFill(start_color="FFC6EFCE", end_color="FFC6EFCE", fill_type="solid")
FILL_RED = PatternFill(start_color="FFFFC7CE", end_color="FFFFC7CE", fill_type="solid")


def build_stress_test(wb: Workbook) -> Worksheet:
    ws = wb.create_sheet("Stress_Test")
    ws.sheet_properties.tabColor = TAB_COLOR_INPUT

    ws["A1"] = "Stress_Test — One-at-a-Time Sensitivity on the Active Scenario"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = (
        "Each row shocks a single driver on the currently active scenario, re-solves both "
        "circularity loops, records the result, then restores the original value before "
        "the next row — never cumulative. Run via the [Batch] Run Stress Test button on Cover."
    )
    ws["A2"].font = Font(italic=True, size=9)

    ws.cell(row=ROW_SPEC_HEADER, column=1, value="Shock Specification").font = Font(bold=True)
    for col, header in enumerate(
        ("Variable", "Assumptions_Constant Row", "Shock Type", "Shock Value"), start=1
    ):
        cell = ws.cell(row=ROW_SPEC_TABLE_HEADER, column=col, value=header)
        cell.font = Font(bold=True)

    for i, (label, row, shock_type, shock_value) in enumerate(SHOCK_SPECS):
        r = ROW_FIRST_SPEC + i
        ws.cell(row=r, column=1, value=label).font = Font(color=COLOR_INPUT)
        row_cell = ws.cell(row=r, column=2, value=row)
        row_cell.font = Font(color=COLOR_INPUT)
        type_cell = ws.cell(row=r, column=3, value=shock_type)
        type_cell.font = Font(color=COLOR_INPUT)
        value_cell = ws.cell(row=r, column=4, value=shock_value)
        value_cell.font = Font(color=COLOR_INPUT)
        value_cell.number_format = "0.00%" if shock_type == "relative" else "0.0000"

    ws.cell(row=ROW_RESULTS_HEADER, column=1,
            value="Results (written as values by the macro — same convention as Batch_Results)").font = Font(
        bold=True)
    for col, header in enumerate(
        ("Variable", "Shock", "Base EIRR", "Stressed EIRR", "Δ EIRR (bps)",
         "Base PIRR", "Stressed PIRR", "Δ PIRR (bps)", "Gearing", "Min DSCR", "Converged?"),
        start=1,
    ):
        cell = ws.cell(row=ROW_RESULTS_TABLE_HEADER, column=col, value=header)
        cell.font = Font(bold=True)

    for i in range(N_SHOCKS):
        r = ROW_FIRST_RESULT + i
        for col, fmt in (
            (COL_BASE_EIRR, "0.00%"), (COL_STRESSED_EIRR, "0.00%"), (COL_DELTA_EIRR_BPS, "#,##0"),
            (COL_BASE_PIRR, "0.00%"), (COL_STRESSED_PIRR, "0.00%"), (COL_DELTA_PIRR_BPS, "#,##0"),
            (COL_GEARING, "0.00%"), (COL_MIN_DSCR, "0.0000"),
        ):
            ws.cell(row=r, column=col).number_format = fmt

    # Conditional formatting: a bad (negative-return-widening) EIRR delta reads red, an
    # improving one reads green — same traffic-light convention as the Dashboard.
    delta_eirr_range = f"E{ROW_FIRST_RESULT}:E{ROW_LAST_RESULT}"
    ws.conditional_formatting.add(delta_eirr_range, FormulaRule(formula=[f"E{ROW_FIRST_RESULT}<0"], fill=FILL_RED))
    ws.conditional_formatting.add(delta_eirr_range, FormulaRule(formula=[f"E{ROW_FIRST_RESULT}>=0"], fill=FILL_GREEN))

    converged_range = f"K{ROW_FIRST_RESULT}:K{ROW_LAST_RESULT}"
    ws.conditional_formatting.add(
        converged_range, FormulaRule(formula=[f'K{ROW_FIRST_RESULT}="NOT CONV"'], fill=FILL_RED))

    ws.cell(row=ROW_META_HEADER, column=1, value="Run Metadata").font = Font(bold=True)
    ws.cell(row=ROW_LAST_RUN, column=1, value="Last Run")
    ws.cell(row=ROW_LAST_RUN, column=2, value="(never)").font = Font(color=COLOR_FORMULA)
    ws.cell(row=ROW_RUN_SCENARIO, column=1, value="Scenario Stress-Tested")
    ws.cell(row=ROW_RUN_SCENARIO, column=2, value="(never)").font = Font(color=COLOR_FORMULA)

    # No row-count named range: the macro loops the spec column until it hits a blank
    # cell, the same convention DrawControlPanel already uses for ButtonSpec — one fewer
    # thing to keep in sync between the build and the VBA.
    _add_named_range(wb, "StressTest_SpecAnchor", "Stress_Test", f"A{ROW_FIRST_SPEC}")
    _add_named_range(wb, "StressTest_ResultsAnchor", "Stress_Test", f"A{ROW_FIRST_RESULT}")
    _add_named_range(wb, "StressTest_LastRun", "Stress_Test", f"B{ROW_LAST_RUN}")
    _add_named_range(wb, "StressTest_RunScenario", "Stress_Test", f"B{ROW_RUN_SCENARIO}")

    ws.column_dimensions["A"].width = 26
    ws.freeze_panes = ws.cell(row=ROW_SPEC_TABLE_HEADER + 1, column=1)

    return ws


def _add_named_range(wb: Workbook, name: str, sheet: str, cell: str) -> None:
    import re

    from openpyxl.workbook.defined_name import DefinedName

    match = re.match(r"([A-Za-z]+)(\d+)", cell)
    col, row = match.group(1), match.group(2)
    wb.defined_names.add(DefinedName(name, attr_text=f"'{sheet}'!${col}${row}"))
