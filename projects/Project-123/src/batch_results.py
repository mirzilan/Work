from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.workbook import Workbook

import assumptions_constant as const
from workbook_builder import COLOR_LINK, TAB_COLOR_OUTPUT

ROW_TITLE = 1
ROW_LAST_RUN = 3
ROW_MODE = 4
ROW_TABLE_HEADER = 6
ROW_FIRST_SCENARIO = 7

# The batch macro writes values into these columns, one row per scenario. Nothing here is
# a formula: a formula would recalculate to the *currently selected* scenario the moment
# the run finished, so every row would end up showing the same numbers.
COLUMNS = [
    ("Scenario", 10, None),
    ("Name", 16, None),
    ("Total Project Cost", 18, "#,##0"),
    ("Debt Facility", 16, "#,##0"),
    ("Implied Gearing", 15, "0.00%"),
    ("Min DSCR", 12, "0.0000"),
    ("Min LLCR", 12, "0.000"),
    ("EIRR", 12, "0.00%"),
    ("PIRR", 12, "0.00%"),
    ("Annual Revenue", 16, "#,##0"),
    ("Solve", 12, None),
    ("Goal Seek", 24, None),
    ("Checks", 16, None),
]


def build_batch_results(wb: Workbook) -> Worksheet:
    ws = wb.create_sheet("Batch_Results")
    ws.sheet_properties.tabColor = TAB_COLOR_OUTPUT

    ws.cell(row=ROW_TITLE, column=1,
            value="Batch_Results — written by Run All 10 Scenarios").font = Font(bold=True, size=12)

    ws.cell(row=ROW_LAST_RUN, column=1, value="Last Batch Run")
    stamp = ws.cell(row=ROW_LAST_RUN, column=2, value="(never)")
    stamp.font = Font(italic=True)

    ws.cell(row=ROW_MODE, column=1, value="Mode Used")
    mode = ws.cell(row=ROW_MODE, column=2, value="(none)")
    mode.font = Font(italic=True)

    for i, (header, width, _fmt) in enumerate(COLUMNS):
        cell = ws.cell(row=ROW_TABLE_HEADER, column=i + 1, value=header)
        cell.font = Font(bold=True)
        ws.column_dimensions[cell.column_letter].width = width

    for s in range(const.N_SCENARIOS):
        row = ROW_FIRST_SCENARIO + s
        ws.cell(row=row, column=1, value=s + 1)

        # The name is the one safe formula here: it reads the scenario's own header cell
        # rather than whatever is currently active, so it stays correct without the macro.
        name_cell = ws.cell(row=row, column=2)
        name_cell.value = (
            f"=Assumptions_Constant!{_scenario_col(s)}${const.ROW_SCENARIO_NAME}"
        )
        name_cell.font = Font(color=COLOR_LINK)

        for i, (_header, _width, fmt) in enumerate(COLUMNS):
            if fmt:
                ws.cell(row=row, column=i + 1).number_format = fmt

    last_row = ROW_FIRST_SCENARIO + const.N_SCENARIOS - 1
    ws.cell(row=last_row + 2, column=1,
            value="Values are written by the macro, not linked — see the note in batch_results.py"
            ).font = Font(italic=True, size=9)

    _name(wb, "BatchResults_Anchor", "Batch_Results", f"A{ROW_FIRST_SCENARIO}")
    _name(wb, "BatchResults_LastRun", "Batch_Results", f"B{ROW_LAST_RUN}")
    _name(wb, "BatchResults_Mode", "Batch_Results", f"B{ROW_MODE}")

    ws.freeze_panes = ws.cell(row=ROW_FIRST_SCENARIO, column=3)

    return ws


def _scenario_col(scenario_index: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(const.FIRST_SCENARIO_COL + scenario_index)


def _name(wb: Workbook, name: str, sheet: str, cell: str) -> None:
    import re

    from openpyxl.workbook.defined_name import DefinedName

    m = re.match(r"([A-Za-z]+)(\d+)", cell)
    wb.defined_names.add(DefinedName(name, attr_text=f"'{sheet}'!${m.group(1)}${m.group(2)}"))
