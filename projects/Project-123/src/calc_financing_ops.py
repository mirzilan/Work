from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.workbook import Workbook

import assumptions_model as model
import cover_refs as refs
from inputs import ProjectInputs
from timeline import Timeline
from workbook_builder import (
    FIRST_DATA_COL,
    COLOR_INPUT,
    COLOR_FORMULA,
    COLOR_LINK,
    TAB_COLOR_CALC,
    col_letter,
)

CELL_SIZING_MODE = "B3"
CELL_INTEREST_RATE = "B4"
CELL_TARGET_DSCR = "B5"
CELL_MAX_GEARING = "B6"
CELL_STAGED_DEBT_SIZE = "B7"
CELL_SCULPTED_CAPACITY = "B8"
CELL_TENOR_YEARS = "B9"
CELL_CONVERGENCE_GAP = "B10"
CELL_IMPLIED_GEARING = "B11"
CELL_DSRA_LC_FEE_RATE = "B12"
CELL_DSRA_METHOD = "B13"

ABS_INTEREST_RATE = "$B$4"
ABS_TARGET_DSCR = "$B$5"
ABS_TENOR_YEARS = "$B$9"
ABS_SIZING_MODE = "$B$3"
ABS_DSRA_LC_FEE_RATE = "$B$12"
ABS_DSRA_METHOD = "$B$13"

ROW_DATE_HEADER = 15
ROW_QUARTER_INDEX = 16

ROW_OPENING_BAL = 18
ROW_INTEREST = 19
ROW_SCULPT_BASIS = 20   # CFADS / Target DSCR, uncapped — the stream the capacity PV discounts
ROW_DEBT_SERVICE = 21   # actual service, floored at interest and capped at amount outstanding
ROW_PRINCIPAL = 22
ROW_CLOSING_BAL = 23
ROW_DSCR = 24

# The DSRA lives here, not on Calc_CFADS: its requirement is forward debt service, and its
# LC fee has to reach Calc_Tax as a deduction. Sourcing the fee from the waterfall sheet
# would put Tax and CFADS in a cycle no staged cell breaks.
ROW_DSRA_TARGET = 26
ROW_DSRA_BALANCE = 27
ROW_DSRA_FUNDING = 28
ROW_DSRA_LC_FEE = 29
ROW_DSRA_CASH_COST = 30

ROW_LLCR = 32
ROW_PLCR = 33

ROW_CHECK_HEADER = 36
ROW_CHECK_FULLY_AMORTIZED = 37
ROW_CHECK_SCULPT_CONVERGED = 38
ROW_CHECK_MIN_DSCR = 39
ROW_CHECK_PRINCIPAL_FLOORED = 40
ROW_CHECK_CAPPED_BY_MAX_GEARING = 41
ROW_CHECK_PLCR_GE_LLCR = 42
ROW_CHECK_MIN_LLCR = 43

DSRA_METHOD_CASH = refs.DSRA_METHOD_CASH


def build_calc_financing_ops(wb: Workbook, timeline: Timeline, inputs: ProjectInputs) -> Worksheet:
    # Local import: Calc_CFADS links back to this sheet's rows, so a module-level import
    # here would close the loop. The row constant is only needed at build time.
    from calc_cfads import ROW_CFADS as CFADS_ROW_CFADS

    ws = wb.create_sheet("Calc_Financing_Ops")
    ws.sheet_properties.tabColor = TAB_COLOR_CALC

    ws["A1"] = "Calc_Financing_Ops — Quarterly Debt Service (DSCR-sculpted, Loop 2), DSRA, LLCR/PLCR"
    ws["A1"].font = Font(bold=True, size=12)

    n_quarters = len(timeline.operations_quarters)
    tenor_quarters = min(inputs.financing.debt_tenor_years * 4, n_quarters)
    dsra_quarters = max(1, round(inputs.reserves.dsra_target_months / 3))
    first_col = col_letter(0)
    last_col = col_letter(n_quarters - 1)
    tenor_end_col = col_letter(tenor_quarters - 1)
    last_cons_col = col_letter(len(timeline.construction_months) - 1)

    ws["A3"] = "Debt Sizing Mode — linked from Cover"
    ws[CELL_SIZING_MODE] = f"=Cover!{refs.ABS_DEBT_SIZING_MODE}"
    ws[CELL_SIZING_MODE].font = Font(color=COLOR_LINK)

    ws["A4"] = "Interest Rate (Annual) — linked from Assumptions_Model"
    ws[CELL_INTEREST_RATE] = f"=Assumptions_Model!$B${model.ROW_INTEREST_RATE}"
    ws[CELL_INTEREST_RATE].font = Font(color=COLOR_LINK)
    ws[CELL_INTEREST_RATE].number_format = "0.00%"

    ws["A5"] = "Target DSCR — linked from Assumptions_Model"
    ws[CELL_TARGET_DSCR] = f"=Assumptions_Model!$B${model.ROW_TARGET_DSCR}"
    ws[CELL_TARGET_DSCR].font = Font(color=COLOR_LINK)
    ws[CELL_TARGET_DSCR].number_format = "0.00x"

    ws["A6"] = "Max Gearing (cap) — linked from Assumptions_Model"
    ws[CELL_MAX_GEARING] = f"=Assumptions_Model!$B${model.ROW_MAX_GEARING}"
    ws[CELL_MAX_GEARING].font = Font(color=COLOR_LINK)
    ws[CELL_MAX_GEARING].number_format = "0.00%"

    ws["A7"] = "Staged Debt Size ($) — VBA-written, breaks Loop 2 circularity"
    ws[CELL_STAGED_DEBT_SIZE] = float(inputs.capex.total_capex * inputs.financing.debt_pct_of_capex)
    ws[CELL_STAGED_DEBT_SIZE].font = Font(color=COLOR_INPUT, bold=True)
    ws[CELL_STAGED_DEBT_SIZE].number_format = "#,##0"

    # With DSCR locked, the balance recursion is linear:
    #   Balance(t+1) = Balance(t) x (1+r) - CFADS(t)/DSCR
    # Setting Balance(T) = 0 and solving gives a closed form -- the sculpted debt capacity
    # is simply the PV of the debt-service stream at the debt rate. No root-find required;
    # the only iteration left is the tax-shield/IDC fixed point the staged cell breaks.
    # The PV must discount the *uncapped* basis. Using the capped actual service instead
    # makes the fixed point degenerate: once the balance reaches zero the service drops to
    # zero, so the PV just reproduces whatever balance it was handed and any starting debt
    # size looks "converged".
    ws["A8"] = "Sculpted Debt Capacity ($) = PV of sculpting basis, capped at Max Gearing"
    ws[CELL_SCULPTED_CAPACITY] = (
        f"=MIN(NPV({ABS_INTEREST_RATE}/4,{first_col}{ROW_SCULPT_BASIS}:{tenor_end_col}{ROW_SCULPT_BASIS}),"
        f"$B$6*Calc_Capex!${last_cons_col}$17)"
    )
    ws[CELL_SCULPTED_CAPACITY].font = Font(color=COLOR_FORMULA, bold=True)
    ws[CELL_SCULPTED_CAPACITY].number_format = "#,##0"

    ws["A9"] = "Debt Tenor (Years) — linked from Assumptions_Model"
    ws[CELL_TENOR_YEARS] = f"=Assumptions_Model!$B${model.ROW_DEBT_TENOR_YEARS}"
    ws[CELL_TENOR_YEARS].font = Font(color=COLOR_LINK)

    ws["A10"] = "Convergence Gap ($) = Capacity - Staged Debt Size"
    ws[CELL_CONVERGENCE_GAP] = "=$B$8-$B$7"
    ws[CELL_CONVERGENCE_GAP].font = Font(color=COLOR_FORMULA, bold=True)
    ws[CELL_CONVERGENCE_GAP].number_format = "#,##0.00"

    ws["A11"] = "Implied Gearing (solved) = Opening Debt / Total Project Cost"
    ws[CELL_IMPLIED_GEARING] = (
        f"=IFERROR({first_col}{ROW_OPENING_BAL}/Calc_Capex!${last_cons_col}$17,0)"
    )
    ws[CELL_IMPLIED_GEARING].font = Font(color=COLOR_FORMULA, bold=True)
    ws[CELL_IMPLIED_GEARING].number_format = "0.00%"

    ws["A12"] = "DSRA LC Fee (% p.a. on requirement) — linked from Assumptions_Model"
    ws[CELL_DSRA_LC_FEE_RATE] = f"=Assumptions_Model!$B${model.ROW_DSRA_LC_FEE}"
    ws[CELL_DSRA_LC_FEE_RATE].font = Font(color=COLOR_LINK)
    ws[CELL_DSRA_LC_FEE_RATE].number_format = "0.00%"

    ws["A13"] = (
        "DSRA Funding Method — linked from Cover (target = "
        f"{inputs.reserves.dsra_target_months} months of forward debt service)"
    )
    ws[CELL_DSRA_METHOD] = f"=Cover!{refs.ABS_DSRA_METHOD}"
    ws[CELL_DSRA_METHOD].font = Font(color=COLOR_LINK)

    _label(ws, ROW_DATE_HEADER, "Period End Date")
    _label(ws, ROW_QUARTER_INDEX, "Operating Quarter #")
    _label(ws, ROW_OPENING_BAL, "Opening Debt Balance ($)")
    _label(ws, ROW_INTEREST, "Interest ($)")
    _label(ws, ROW_SCULPT_BASIS, "Sculpting Basis ($) = CFADS / Target DSCR (uncapped)")
    _label(ws, ROW_DEBT_SERVICE, "Debt Service ($) — sculpted to Target DSCR")
    _label(ws, ROW_PRINCIPAL, "Principal ($)")
    _label(ws, ROW_CLOSING_BAL, "Closing Debt Balance ($)")
    _label(ws, ROW_DSCR, "DSCR (achieved)")
    _label(ws, ROW_DSRA_TARGET, f"DSRA Target ($) = next {dsra_quarters} quarters' debt service")
    _label(ws, ROW_DSRA_BALANCE, "DSRA Balance ($) — zero when LC-backed")
    _label(ws, ROW_DSRA_FUNDING, "DSRA Funding/(Release) ($) — cash trapped this quarter")
    _label(ws, ROW_DSRA_LC_FEE, "DSRA LC Fee ($) — charged only when LC-backed")
    _label(ws, ROW_DSRA_CASH_COST, "DSRA Cash Cost ($) = funding (cash-funded) or LC fee (LC-backed)")
    _label(ws, ROW_LLCR, "LLCR = (PV of CFADS to loan maturity + DSRA) / debt outstanding")
    _label(ws, ROW_PLCR, "PLCR = (PV of CFADS to end of project life + DSRA) / debt outstanding")

    ws.cell(row=ROW_CHECK_HEADER, column=1, value="Checks").font = Font(bold=True)
    _label(ws, ROW_CHECK_FULLY_AMORTIZED, "Check: Closing Balance = 0 at Debt Tenor End")
    _label(ws, ROW_CHECK_SCULPT_CONVERGED, "Check: Debt Sizing Converged (Loop 2)")
    _label(ws, ROW_CHECK_MIN_DSCR, "Check: Min DSCR over tenor >= Target DSCR")
    _label(ws, ROW_CHECK_PRINCIPAL_FLOORED, "Informational: # quarters principal floored at zero")
    _label(ws, ROW_CHECK_CAPPED_BY_MAX_GEARING, "Informational: debt sizing capped by Max Gearing (1 = capped)")
    _label(ws, ROW_CHECK_PLCR_GE_LLCR, "Check: PLCR >= LLCR in every quarter")
    _label(ws, ROW_CHECK_MIN_LLCR, "Check: LLCR >= 1.00 wherever debt is outstanding")

    pmt_formula = (
        f"-PMT({ABS_INTEREST_RATE}/4,{ABS_TENOR_YEARS}*4,"
        f"Calc_Financing_Cons!{last_cons_col}25)"
    )

    for i, period in enumerate(timeline.operations_quarters):
        col = col_letter(i)
        prev = col_letter(i - 1) if i > 0 else None

        ws[f"{col}{ROW_DATE_HEADER}"] = period.end
        ws[f"{col}{ROW_DATE_HEADER}"].number_format = "mmm-yy"
        ws[f"{col}{ROW_QUARTER_INDEX}"] = i + 1

        if i == 0:
            _link(ws, col, ROW_OPENING_BAL, f"Calc_Financing_Cons!{last_cons_col}25")
        else:
            _formula(ws, col, ROW_OPENING_BAL, f"={prev}{ROW_CLOSING_BAL}")

        _formula(ws, col, ROW_INTEREST, f"={col}{ROW_OPENING_BAL}*{ABS_INTEREST_RATE}/4")

        if i < tenor_quarters:
            _formula(ws, col, ROW_SCULPT_BASIS, f"=Calc_CFADS!{col}{CFADS_ROW_CFADS}/{ABS_TARGET_DSCR}")
            # Actual service: never below interest (which would capitalise unpaid interest),
            # never above what is still outstanding.
            sculpted = (
                f"MAX({col}{ROW_INTEREST},MIN({col}{ROW_SCULPT_BASIS},"
                f"{col}{ROW_OPENING_BAL}+{col}{ROW_INTEREST}))"
            )
            level = f"MIN({pmt_formula},{col}{ROW_OPENING_BAL}+{col}{ROW_INTEREST})"
            _formula(ws, col, ROW_DEBT_SERVICE, f'=IF({ABS_SIZING_MODE}="{refs.SCULPTED}",{sculpted},{level})')
        else:
            _formula(ws, col, ROW_SCULPT_BASIS, "=0")
            _formula(ws, col, ROW_DEBT_SERVICE, "=0")

        _formula(ws, col, ROW_PRINCIPAL, f"={col}{ROW_DEBT_SERVICE}-{col}{ROW_INTEREST}")
        _formula(ws, col, ROW_CLOSING_BAL, f"={col}{ROW_OPENING_BAL}-{col}{ROW_PRINCIPAL}")

        dscr_cell = ws[f"{col}{ROW_DSCR}"]
        dscr_cell.value = (
            f'=IF({col}{ROW_DEBT_SERVICE}=0,"",'
            f"Calc_CFADS!{col}{CFADS_ROW_CFADS}/{col}{ROW_DEBT_SERVICE})"
        )
        dscr_cell.font = Font(color=COLOR_LINK)
        dscr_cell.number_format = "0.00x"

        _write_dsra(ws, i, col, prev, n_quarters, dsra_quarters)
        _write_coverage_ratios(ws, i, col, n_quarters, tenor_quarters,
                               tenor_end_col, last_col, CFADS_ROW_CFADS)

    # A debt size converged to within the sizing tolerance leaves a tenor-end residual that
    # compounds over the tenor, so an absolute epsilon would false-fail on longer tenors.
    # Judge it on materiality instead: 0.001% of the opening balance, floored at the tolerance.
    _check(ws, last_col, ROW_CHECK_FULLY_AMORTIZED,
           f"=IF(ABS({tenor_end_col}{ROW_CLOSING_BAL})"
           f"<=MAX(Cover!{refs.ABS_DEBT_SIZING_TOLERANCE},{first_col}{ROW_OPENING_BAL}*0.00001),1,0)")
    _check(ws, last_col, ROW_CHECK_SCULPT_CONVERGED,
           f'=IF({ABS_SIZING_MODE}<>"{refs.SCULPTED}",1,'
           f"IF(ABS($B$10)<=Cover!{refs.ABS_DEBT_SIZING_TOLERANCE},1,0))")
    _check(ws, last_col, ROW_CHECK_MIN_DSCR,
           f"=IF(MIN({first_col}{ROW_DSCR}:{tenor_end_col}{ROW_DSCR})>={ABS_TARGET_DSCR}-0.001,1,0)")
    _check(ws, last_col, ROW_CHECK_PRINCIPAL_FLOORED,
           f"=SUMPRODUCT(--({first_col}{ROW_SCULPT_BASIS}:{tenor_end_col}{ROW_SCULPT_BASIS}"
           f"<{first_col}{ROW_INTEREST}:{tenor_end_col}{ROW_INTEREST}))")

    # Surfaces when policy, not cash flow, is the binding constraint — the project could
    # carry more debt than Max Gearing permits.
    _check(ws, last_col, ROW_CHECK_CAPPED_BY_MAX_GEARING,
           f"=IF(NPV({ABS_INTEREST_RATE}/4,{first_col}{ROW_SCULPT_BASIS}:{tenor_end_col}{ROW_SCULPT_BASIS})"
           f">$B$6*Calc_Capex!${last_cons_col}$17,1,0)")

    # PLCR discounts CFADS over a strictly longer horizon than LLCR off the same balance,
    # so it can never be the smaller of the two — if it is, one of the ranges is wrong.
    _check(ws, last_col, ROW_CHECK_PLCR_GE_LLCR,
           f"=IF(SUMPRODUCT(--({first_col}{ROW_LLCR}:{last_col}{ROW_LLCR}"
           f">{first_col}{ROW_PLCR}:{last_col}{ROW_PLCR}+0.000001))=0,1,0)")

    # Both ratios read 0 where no debt is outstanding, so the test only bites on live debt.
    _check(ws, last_col, ROW_CHECK_MIN_LLCR,
           f"=IF(SUMPRODUCT(({first_col}{ROW_LLCR}:{last_col}{ROW_LLCR}>0)*"
           f"({first_col}{ROW_LLCR}:{last_col}{ROW_LLCR}<1))=0,1,0)")

    _add_named_range(wb, "StagedDebtSize", "Calc_Financing_Ops", CELL_STAGED_DEBT_SIZE)
    _add_named_range(wb, "SculptedDebtCapacity", "Calc_Financing_Ops", CELL_SCULPTED_CAPACITY)
    _add_named_range(wb, "DebtSizeConvergenceGap", "Calc_Financing_Ops", CELL_CONVERGENCE_GAP)
    _add_named_range(wb, "ImpliedGearing", "Calc_Financing_Ops", CELL_IMPLIED_GEARING)
    _add_named_range(wb, "FinOps_LastCol", "Calc_Financing_Ops", f"{last_col}1")
    _add_named_range(wb, "FinOps_TenorEndCol", "Calc_Financing_Ops", f"{tenor_end_col}1")

    ws.freeze_panes = ws.cell(row=ROW_PLCR + 1, column=FIRST_DATA_COL)
    ws.column_dimensions["A"].width = 58

    return ws


def _write_dsra(ws: Worksheet, i: int, col: str, prev: str | None,
                n_quarters: int, dsra_quarters: int) -> None:
    """The requirement is forward-looking, so it winds itself down to zero as the tail of
    the debt service stream runs out — no explicit release event needed."""
    if i + 1 < n_quarters:
        window_start = col_letter(i + 1)
        window_end = col_letter(min(i + dsra_quarters, n_quarters - 1))
        _formula(ws, col, ROW_DSRA_TARGET,
                 f"=SUM({window_start}{ROW_DEBT_SERVICE}:{window_end}{ROW_DEBT_SERVICE})")
    else:
        _formula(ws, col, ROW_DSRA_TARGET, "=0")

    _formula(ws, col, ROW_DSRA_BALANCE,
             f'=IF({ABS_DSRA_METHOD}="{DSRA_METHOD_CASH}",{col}{ROW_DSRA_TARGET},0)')

    if prev is None:
        _formula(ws, col, ROW_DSRA_FUNDING, f"={col}{ROW_DSRA_BALANCE}")
    else:
        _formula(ws, col, ROW_DSRA_FUNDING, f"={col}{ROW_DSRA_BALANCE}-{prev}{ROW_DSRA_BALANCE}")

    _formula(ws, col, ROW_DSRA_LC_FEE,
             f'=IF({ABS_DSRA_METHOD}="{DSRA_METHOD_CASH}",0,'
             f"{col}{ROW_DSRA_TARGET}*{ABS_DSRA_LC_FEE_RATE}/4)")

    # Exactly one of the two is non-zero for any given method, so adding them keeps the
    # waterfall to a single line without an extra IF.
    _formula(ws, col, ROW_DSRA_CASH_COST, f"={col}{ROW_DSRA_FUNDING}+{col}{ROW_DSRA_LC_FEE}")


def _write_coverage_ratios(ws: Worksheet, i: int, col: str, n_quarters: int,
                           tenor_quarters: int, tenor_end_col: str, last_col: str,
                           cfads_row: int) -> None:
    """Forward PV of CFADS over debt outstanding. Both read 0 once the debt is repaid, so
    the rows stay usable in MIN/SUMPRODUCT checks without ISNUMBER gymnastics."""
    balance = f"{col}{ROW_CLOSING_BAL}"
    dsra = f"{col}{ROW_DSRA_BALANCE}"

    for row, horizon_end_col, horizon_last_index in (
        (ROW_LLCR, tenor_end_col, tenor_quarters - 1),
        (ROW_PLCR, last_col, n_quarters - 1),
    ):
        cell = ws[f"{col}{row}"]
        if i >= horizon_last_index:
            cell.value = "=0"
        else:
            fwd_start = col_letter(i + 1)
            cell.value = (
                f"=IF({balance}<=1,0,"
                f"(NPV({ABS_INTEREST_RATE}/4,Calc_CFADS!{fwd_start}{cfads_row}:"
                f"Calc_CFADS!{horizon_end_col}{cfads_row})+{dsra})/{balance})"
            )
        cell.font = Font(color=COLOR_FORMULA)
        cell.number_format = "0.00x"


def _label(ws: Worksheet, row: int, label: str) -> None:
    ws.cell(row=row, column=1, value=label)


def _link(ws: Worksheet, col: str, row: int, formula: str) -> None:
    cell = ws[f"{col}{row}"]
    cell.value = f"={formula}"
    cell.font = Font(color=COLOR_LINK)
    cell.number_format = "#,##0"


def _formula(ws: Worksheet, col: str, row: int, formula: str) -> None:
    cell = ws[f"{col}{row}"]
    cell.value = formula
    cell.font = Font(color=COLOR_FORMULA)
    cell.number_format = "#,##0"


def _check(ws: Worksheet, col: str, row: int, formula: str) -> None:
    cell = ws[f"{col}{row}"]
    cell.value = formula
    cell.font = Font(color=COLOR_FORMULA)


def _add_named_range(wb: Workbook, name: str, sheet: str, cell: str) -> None:
    import re

    from openpyxl.workbook.defined_name import DefinedName

    match = re.match(r"([A-Za-z]+)(\d+)", cell)
    col, row = match.group(1), match.group(2)
    wb.defined_names.add(DefinedName(name, attr_text=f"'{sheet}'!${col}${row}"))
