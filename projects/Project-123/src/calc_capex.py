from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.workbook import Workbook

from inputs import ProjectInputs
from timeline import Timeline
from workbook_builder import (
    FIRST_DATA_COL,
    COLOR_INPUT,
    COLOR_FORMULA,
    COLOR_LINK,
    TAB_COLOR_CALC,
    col_letter,
)

ROW_DATE_HEADER = 2
ROW_MONTH_INDEX = 3
ROW_PHASING_PCT = 5
ROW_CAPEX_DRAW = 6
ROW_CUM_CAPEX_DRAW = 7
ROW_DEBT_FUNDING_PCT = 9
ROW_DEBT_DRAW = 10
ROW_EQUITY_DRAW = 11
ROW_CUM_DEBT_DRAW = 12
ROW_CUM_EQUITY_DRAW = 13
ROW_IDC_CAPITALIZED = 15
ROW_CUM_IDC = 16
ROW_TOTAL_PROJECT_COST = 17

ROW_CHECK_HEADER = 20
ROW_CHECK_FUNDING_TIES = 21
ROW_CHECK_TOTAL_MATCHES_INPUT = 22


def build_calc_capex(wb: Workbook, timeline: Timeline, inputs: ProjectInputs) -> Worksheet:
    ws = wb.create_sheet("Calc_Capex")
    ws.sheet_properties.tabColor = TAB_COLOR_CALC

    ws["A1"] = "Calc_Capex — Monthly Construction Capex & Funding"
    ws["A1"].font = Font(bold=True, size=12)

    _write_row_label(ws, ROW_DATE_HEADER, "Period End Date")
    _write_row_label(ws, ROW_MONTH_INDEX, "Construction Month #")
    _write_row_label(ws, ROW_PHASING_PCT, "Capex Phasing %")
    _write_row_label(ws, ROW_CAPEX_DRAW, "Capex Draw ($)")
    _write_row_label(ws, ROW_CUM_CAPEX_DRAW, "Cumulative Capex Draw ($)")
    _write_row_label(ws, ROW_DEBT_FUNDING_PCT, "Gearing (Debt % of Total Project Cost)")
    _write_row_label(ws, ROW_DEBT_DRAW, "Debt Draw ($) — linked from Calc_Financing_Cons")
    _write_row_label(ws, ROW_EQUITY_DRAW, "Equity Draw ($) — linked from Calc_Financing_Cons")
    _write_row_label(ws, ROW_CUM_DEBT_DRAW, "Cumulative Debt Draw ($) — linked from Calc_Financing_Cons")
    _write_row_label(ws, ROW_CUM_EQUITY_DRAW, "Cumulative Equity Draw ($) — linked from Calc_Financing_Cons")
    _write_row_label(ws, ROW_IDC_CAPITALIZED, "IDC ($, monthly) — solved via Loop 1, linked from Calc_Financing_Cons")
    _write_row_label(ws, ROW_CUM_IDC, "Cumulative IDC ($)")
    _write_row_label(ws, ROW_TOTAL_PROJECT_COST, "Total Project Cost (Cumulative) = Cum Capex + Cum IDC")

    ws.cell(row=ROW_CHECK_HEADER, column=1, value="Checks").font = Font(bold=True)
    _write_row_label(ws, ROW_CHECK_FUNDING_TIES, "Check: Cum Debt + Cum Equity = Cum Capex + Cum IDC")
    _write_row_label(ws, ROW_CHECK_TOTAL_MATCHES_INPUT, "Check: Final Cumulative Capex = Total Capex Input")

    ws["A4"] = "Total Capex Input ($) — linked from Assumptions_Model"
    ws["B4"] = "=Assumptions_Model!$B$3"
    ws["B4"].font = Font(color=COLOR_LINK)
    ws["B4"].number_format = "#,##0"
    wb.defined_names.add(_named_range("TotalCapex", "Calc_Capex", "B4"))

    ws["A9"] = "Debt Funding % — linked from Assumptions_Model (Stage 1a: fixed ratio; Stage 1b: Loop 1 draw sequencing)"
    ws["B9"] = "=Assumptions_Model!$B$4"
    ws["B9"].font = Font(color=COLOR_LINK)
    ws["B9"].number_format = "0.00%"

    n_months = len(timeline.construction_months)

    for i, period in enumerate(timeline.construction_months):
        col = col_letter(i)

        ws[f"{col}{ROW_DATE_HEADER}"] = period.end
        ws[f"{col}{ROW_DATE_HEADER}"].number_format = "mmm-yy"

        ws[f"{col}{ROW_MONTH_INDEX}"] = i + 1

        # Phasing % — green link from Assumptions_Model (same column alignment)
        phasing_cell = ws[f"{col}{ROW_PHASING_PCT}"]
        phasing_cell.value = f"=Assumptions_Model!{col}16"  # ROW_PHASING_PCT in assumptions_model.py
        phasing_cell.font = Font(color=COLOR_LINK)
        phasing_cell.number_format = "0.00%"

        # Capex Draw = phasing % * total capex — black formula
        draw_cell = ws[f"{col}{ROW_CAPEX_DRAW}"]
        draw_cell.value = f"={col}{ROW_PHASING_PCT}*$B$4"
        draw_cell.font = Font(color=COLOR_FORMULA)
        draw_cell.number_format = "#,##0"

        # Cumulative capex draw
        cum_cell = ws[f"{col}{ROW_CUM_CAPEX_DRAW}"]
        if i == 0:
            cum_cell.value = f"={col}{ROW_CAPEX_DRAW}"
        else:
            prev_col = col_letter(i - 1)
            cum_cell.value = f"={prev_col}{ROW_CUM_CAPEX_DRAW}+{col}{ROW_CAPEX_DRAW}"
        cum_cell.font = Font(color=COLOR_FORMULA)
        cum_cell.number_format = "#,##0"

        # Funding is owned by Calc_Financing_Cons (drawdown method + IDC solve live there);
        # these rows are presentation links so the capex sheet reads as a full sources/uses view.
        for row, source_row in (
            (ROW_DEBT_DRAW, 21),
            (ROW_EQUITY_DRAW, 22),
            (ROW_CUM_DEBT_DRAW, 23),
            (ROW_CUM_EQUITY_DRAW, 24),
            (ROW_IDC_CAPITALIZED, 18),
        ):
            cell = ws[f"{col}{row}"]
            cell.value = f"=Calc_Financing_Cons!{col}{source_row}"
            cell.font = Font(color=COLOR_LINK)
            cell.number_format = "#,##0"

        cum_idc_cell = ws[f"{col}{ROW_CUM_IDC}"]
        if i == 0:
            cum_idc_cell.value = f"={col}{ROW_IDC_CAPITALIZED}"
        else:
            prev_col = col_letter(i - 1)
            cum_idc_cell.value = f"={prev_col}{ROW_CUM_IDC}+{col}{ROW_IDC_CAPITALIZED}"
        cum_idc_cell.font = Font(color=COLOR_FORMULA)
        cum_idc_cell.number_format = "#,##0"

        total_cost_cell = ws[f"{col}{ROW_TOTAL_PROJECT_COST}"]
        total_cost_cell.value = f"={col}{ROW_CUM_CAPEX_DRAW}+{col}{ROW_CUM_IDC}"
        total_cost_cell.font = Font(color=COLOR_FORMULA)
        total_cost_cell.number_format = "#,##0"

    last_col = col_letter(n_months - 1)

    # Checks
    check_ties = ws[f"{last_col}{ROW_CHECK_FUNDING_TIES}"]
    check_ties.value = (
        f"=IF(ROUND({last_col}{ROW_CUM_DEBT_DRAW}+{last_col}{ROW_CUM_EQUITY_DRAW}"
        f"-{last_col}{ROW_CUM_CAPEX_DRAW}-{last_col}{ROW_CUM_IDC},2)=0,1,0)"
    )
    check_ties.font = Font(color=COLOR_FORMULA)

    check_total = ws[f"{last_col}{ROW_CHECK_TOTAL_MATCHES_INPUT}"]
    check_total.value = f"=IF(ROUND({last_col}{ROW_CUM_CAPEX_DRAW}-$B$4,2)=0,1,0)"
    check_total.font = Font(color=COLOR_FORMULA)

    wb.defined_names.add(_named_range("Capex_FundingTiesCheck", "Calc_Capex", f"{last_col}{ROW_CHECK_FUNDING_TIES}"))
    wb.defined_names.add(_named_range("Capex_TotalMatchesInputCheck", "Calc_Capex", f"{last_col}{ROW_CHECK_TOTAL_MATCHES_INPUT}"))
    wb.defined_names.add(_named_range("Capex_LastCol", "Calc_Capex", f"{last_col}1"))
    wb.defined_names.add(_named_range("Capex_CumTotal_Last", "Calc_Capex", f"{last_col}{ROW_CUM_CAPEX_DRAW}"))
    wb.defined_names.add(_named_range("Capex_TotalProjectCost_Last", "Calc_Capex", f"{last_col}{ROW_TOTAL_PROJECT_COST}"))

    ws.freeze_panes = ws.cell(row=ROW_TOTAL_PROJECT_COST + 1, column=FIRST_DATA_COL)

    return ws


def _write_row_label(ws: Worksheet, row: int, label: str) -> None:
    ws.cell(row=row, column=1, value=label)


def _named_range(name: str, sheet: str, cell: str):
    import re

    from openpyxl.workbook.defined_name import DefinedName

    match = re.match(r"([A-Za-z]+)(\d+)", cell)
    col, row = match.group(1), match.group(2)
    return DefinedName(name, attr_text=f"'{sheet}'!${col}${row}")
