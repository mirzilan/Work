from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.workbook import Workbook

import assumptions_constant as const
import cover_refs as refs
from inputs import ProjectInputs
from timeline import Timeline
from workbook_builder import (
    FIRST_DATA_COL,
    COLOR_INPUT,
    COLOR_FORMULA,
    COLOR_LINK,
    TAB_COLOR_INPUT,
    col_letter,
)

N_SCENARIOS = 10
N_ESC_FACTORS = 4

ROW_DATE_HEADER = 4
ROW_QUARTER_INDEX = 5

ROW_REV_BLOCK_TITLE = 7
ROW_REV_FIRST_SCENARIO = 8
ROW_REV_ACTIVE = ROW_REV_FIRST_SCENARIO + N_SCENARIOS          # 18

ROW_OPEX_BLOCK_TITLE = 20
ROW_OPEX_FIRST_SCENARIO = 21
ROW_OPEX_ACTIVE = ROW_OPEX_FIRST_SCENARIO + N_SCENARIOS        # 31

# Lumpy/overhaul maintenance capex in dollars. Routine maintenance is a % of revenue and
# lives on Assumptions_Constant — only the irregular, date-specific spend needs a profile.
ROW_MAINT_BLOCK_TITLE = 33
ROW_MAINT_FIRST_SCENARIO = 34
ROW_MAINT_ACTIVE = ROW_MAINT_FIRST_SCENARIO + N_SCENARIOS      # 44

ROW_ESC_TITLE = 47
ROW_ESC_HEADER = 48
ROW_ESC_FIRST_FACTOR = 49                                      # factors 1-4 -> rows 49-52
ROW_ESC_LAST_FACTOR = ROW_ESC_FIRST_FACTOR + N_ESC_FACTORS - 1

ROW_REV_ESC_SELECTOR = 54
ROW_OPEX_ESC_SELECTOR = 55
ROW_REV_ESC_ACTIVE = 57
ROW_OPEX_ESC_ACTIVE = 58

ROW_CHECK_HEADER = 61
ROW_CHECK_ESC_BASE = 62
ROW_CHECK_VOLUME_POSITIVE = 63
ROW_CHECK_MAINT_NONNEG = 64

# Escalation rates live on Assumptions_Constant so they vary by scenario; only the index
# derivation belongs here, at this sheet's own resolution.
CONST_ESC_FACTOR_FIRST_ROW = const.ROW_ESC_FACTOR_1


def build_assumptions_periodic_ops(wb: Workbook, timeline: Timeline, inputs: ProjectInputs) -> Worksheet:
    ws = wb.create_sheet("Assumptions_Periodic_Ops")
    ws.sheet_properties.tabColor = TAB_COLOR_INPUT

    ws["A1"] = "Assumptions_Periodic_Ops — Quarterly, Scenarios as Row Blocks"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = (
        "Volume indices default to 1.00 (flat). Escalation rates are scenario-varying and "
        "live on Assumptions_Constant; this sheet compounds them to a quarterly index."
    )
    ws["A2"].font = Font(italic=True, size=9)

    n_q = len(timeline.operations_quarters)
    first_col = col_letter(0)
    last_col = col_letter(n_q - 1)

    ws.cell(row=ROW_DATE_HEADER, column=1, value="Period End Date")
    ws.cell(row=ROW_QUARTER_INDEX, column=1, value="Operating Quarter #")

    for i, period in enumerate(timeline.operations_quarters):
        col = col_letter(i)
        date_cell = ws[f"{col}{ROW_DATE_HEADER}"]
        date_cell.value = period.end
        date_cell.number_format = "mmm-yy"
        ws[f"{col}{ROW_QUARTER_INDEX}"] = i + 1

    _volume_block(ws, n_q, ROW_REV_BLOCK_TITLE, ROW_REV_FIRST_SCENARIO, ROW_REV_ACTIVE,
                  "Revenue Volume Index (by scenario)", "ACTIVE — Revenue Volume Index")
    _volume_block(ws, n_q, ROW_OPEX_BLOCK_TITLE, ROW_OPEX_FIRST_SCENARIO, ROW_OPEX_ACTIVE,
                  "Opex Volume Index (by scenario)", "ACTIVE — Opex Volume Index")

    _maintenance_block(ws, n_q, inputs)
    _escalation_library(ws, n_q)

    ws.cell(row=ROW_CHECK_HEADER, column=1, value="Checks").font = Font(bold=True)
    ws.cell(row=ROW_CHECK_ESC_BASE, column=1, value="Check: All escalation indices start at 1.00")
    ws.cell(row=ROW_CHECK_VOLUME_POSITIVE, column=1, value="Check: Active volume indices > 0 in every quarter")
    ws.cell(row=ROW_CHECK_MAINT_NONNEG, column=1, value="Check: Active lumpy maintenance capex >= 0 in every quarter")

    maint_nonneg = ws.cell(row=ROW_CHECK_MAINT_NONNEG, column=2)
    maint_nonneg.value = f"=IF(MIN({first_col}{ROW_MAINT_ACTIVE}:{last_col}{ROW_MAINT_ACTIVE})>=0,1,0)"
    maint_nonneg.font = Font(color=COLOR_FORMULA)

    esc_base = ws.cell(row=ROW_CHECK_ESC_BASE, column=2)
    esc_base.value = (
        f"=IF(SUMPRODUCT(--(ROUND({first_col}{ROW_ESC_FIRST_FACTOR}:{first_col}{ROW_ESC_LAST_FACTOR}-1,9)<>0))=0,1,0)"
    )
    esc_base.font = Font(color=COLOR_FORMULA)

    vol_positive = ws.cell(row=ROW_CHECK_VOLUME_POSITIVE, column=2)
    vol_positive.value = (
        f"=IF(AND(MIN({first_col}{ROW_REV_ACTIVE}:{last_col}{ROW_REV_ACTIVE})>0,"
        f"MIN({first_col}{ROW_OPEX_ACTIVE}:{last_col}{ROW_OPEX_ACTIVE})>0),1,0)"
    )
    vol_positive.font = Font(color=COLOR_FORMULA)

    _name(wb, "PerOps_EscBaseCheck", "Assumptions_Periodic_Ops", f"B{ROW_CHECK_ESC_BASE}")
    _name(wb, "PerOps_VolumePositiveCheck", "Assumptions_Periodic_Ops", f"B{ROW_CHECK_VOLUME_POSITIVE}")
    _name(wb, "PerOps_MaintNonNegCheck", "Assumptions_Periodic_Ops", f"B{ROW_CHECK_MAINT_NONNEG}")

    ws.column_dimensions["A"].width = 44
    ws.freeze_panes = ws.cell(row=ROW_REV_BLOCK_TITLE + 1, column=FIRST_DATA_COL)

    return ws


def _volume_block(ws: Worksheet, n_q: int, title_row: int, first_row: int, active_row: int,
                  title: str, active_label: str) -> None:
    ws.cell(row=title_row, column=1, value=title).font = Font(bold=True)
    ws.cell(row=active_row, column=1, value=active_label).font = Font(bold=True)

    for s in range(N_SCENARIOS):
        row = first_row + s
        ws.cell(row=row, column=1, value=f"Scenario {s + 1}")
        for i in range(n_q):
            cell = ws[f"{col_letter(i)}{row}"]
            cell.value = 1.0
            cell.font = Font(color=COLOR_INPUT)
            cell.number_format = "0.000"

    for i in range(n_q):
        col = col_letter(i)
        cell = ws[f"{col}{active_row}"]
        cell.value = f"=INDEX({col}${first_row}:{col}${first_row + N_SCENARIOS - 1},Cover!{refs.ABS_ACTIVE_SCENARIO})"
        cell.font = Font(color=COLOR_FORMULA)
        cell.number_format = "0.000"


def _maintenance_block(ws: Worksheet, n_q: int, inputs: ProjectInputs) -> None:
    ws.cell(row=ROW_MAINT_BLOCK_TITLE, column=1,
            value="Lumpy Maintenance Capex ($, by scenario) — overhauls; routine spend is a % of revenue").font = Font(bold=True)
    ws.cell(row=ROW_MAINT_ACTIVE, column=1, value="ACTIVE — Lumpy Maintenance Capex ($)").font = Font(bold=True)

    every = inputs.maintenance.lumpy_overhaul_every_n_quarters
    amount = inputs.maintenance.lumpy_overhaul_amount

    for s in range(N_SCENARIOS):
        row = ROW_MAINT_FIRST_SCENARIO + s
        ws.cell(row=row, column=1, value=f"Scenario {s + 1}")
        for i in range(n_q):
            cell = ws[f"{col_letter(i)}{row}"]
            cell.value = amount if every > 0 and (i + 1) % every == 0 else 0.0
            cell.font = Font(color=COLOR_INPUT)
            cell.number_format = "#,##0"

    for i in range(n_q):
        col = col_letter(i)
        cell = ws[f"{col}{ROW_MAINT_ACTIVE}"]
        cell.value = (
            f"=INDEX({col}${ROW_MAINT_FIRST_SCENARIO}:"
            f"{col}${ROW_MAINT_FIRST_SCENARIO + N_SCENARIOS - 1},Cover!{refs.ABS_ACTIVE_SCENARIO})"
        )
        cell.font = Font(color=COLOR_FORMULA)
        cell.number_format = "#,##0"


def _escalation_library(ws: Worksheet, n_q: int) -> None:
    ws.cell(row=ROW_ESC_TITLE, column=1,
            value="Escalation Library — annual % (active scenario) compounded to a quarterly index").font = Font(bold=True)
    ws.cell(row=ROW_ESC_HEADER, column=1, value="Factor").font = Font(bold=True)
    ws.cell(row=ROW_ESC_HEADER, column=2, value="Annual %").font = Font(bold=True)

    for f in range(N_ESC_FACTORS):
        row = ROW_ESC_FIRST_FACTOR + f
        ws.cell(row=row, column=1, value=f"Escalation Factor {f + 1}")

        rate = ws.cell(row=row, column=2)
        rate.value = f"=Assumptions_Constant!$B${CONST_ESC_FACTOR_FIRST_ROW + f}"
        rate.font = Font(color=COLOR_LINK)
        rate.number_format = "0.00%"

        for i in range(n_q):
            col = col_letter(i)
            cell = ws[f"{col}{row}"]
            if i == 0:
                # Quarter 1 is the base period, so the index starts at 1.00 and escalation
                # accrues from there — within-year compounding, not a year-start step.
                cell.value = 1.0
                cell.font = Font(color=COLOR_FORMULA)
            else:
                prev = col_letter(i - 1)
                cell.value = f"={prev}{row}*(1+$B${row})^(1/4)"
                cell.font = Font(color=COLOR_FORMULA)
            cell.number_format = "0.0000"

    ws.cell(row=ROW_REV_ESC_SELECTOR, column=1, value="Revenue escalation factor # (1-4)")
    rev_sel = ws.cell(row=ROW_REV_ESC_SELECTOR, column=2)
    rev_sel.value = 1
    rev_sel.font = Font(color=COLOR_INPUT)

    ws.cell(row=ROW_OPEX_ESC_SELECTOR, column=1, value="Opex escalation factor # (1-4)")
    opex_sel = ws.cell(row=ROW_OPEX_ESC_SELECTOR, column=2)
    opex_sel.value = 2
    opex_sel.font = Font(color=COLOR_INPUT)

    ws.cell(row=ROW_REV_ESC_ACTIVE, column=1, value="ACTIVE — Revenue Escalation Index").font = Font(bold=True)
    ws.cell(row=ROW_OPEX_ESC_ACTIVE, column=1, value="ACTIVE — Opex Escalation Index").font = Font(bold=True)

    for i in range(n_q):
        col = col_letter(i)
        for active_row, selector_row in (
            (ROW_REV_ESC_ACTIVE, ROW_REV_ESC_SELECTOR),
            (ROW_OPEX_ESC_ACTIVE, ROW_OPEX_ESC_SELECTOR),
        ):
            cell = ws[f"{col}{active_row}"]
            cell.value = (
                f"=INDEX({col}${ROW_ESC_FIRST_FACTOR}:{col}${ROW_ESC_LAST_FACTOR},$B${selector_row})"
            )
            cell.font = Font(color=COLOR_FORMULA)
            cell.number_format = "0.0000"


def _name(wb: Workbook, name: str, sheet: str, cell: str) -> None:
    import re

    from openpyxl.workbook.defined_name import DefinedName

    m = re.match(r"([A-Za-z]+)(\d+)", cell)
    wb.defined_names.add(DefinedName(name, attr_text=f"'{sheet}'!${m.group(1)}${m.group(2)}"))
