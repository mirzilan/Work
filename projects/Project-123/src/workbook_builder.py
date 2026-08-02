import openpyxl
from openpyxl.workbook import Workbook
from openpyxl.styles import Border, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

from inputs import ProjectInputs
from timeline import Timeline

# Column layout, EY/FAST-style: A=Label, B=Units, C=Total (where applicable),
# D=Check (where applicable), E=Remarks, F=buffer (kept clear in case a first-period
# formula ever needs to differ from the rest without disturbing the period grid),
# period data starts at G. Every sheet in the workbook shares this same column scheme —
# FIRST_DATA_COL is the one place that fact lives, so a period index always lands on the
# same column letter on every sheet, which is what makes cross-sheet formulas
# (e.g. `Calc_Revenue_Opex!{col}{row}`) safe to write without re-deriving an offset.
COL_LABEL = 1
COL_UNITS = 2
COL_TOTAL = 3
COL_CHECK = 4
COL_REMARKS = 5
COL_BUFFER = 6
FIRST_DATA_COL = 7  # column G

# FAST/Corality-style font colors (cell-level, per CLAUDE.md convention)
COLOR_INPUT = "FF0000FF"  # blue: hardcoded inputs/assumptions
COLOR_FORMULA = "FF000000"  # black: formulas within sheet
COLOR_LINK = "FF008000"  # green: links from another sheet

# Sheet tab colors (Cover/Assumptions=blue, Calc=grey, Output=green, Check=red)
TAB_COLOR_INPUT = "0000FF"
TAB_COLOR_CALC = "808080"
TAB_COLOR_OUTPUT = "00B050"
TAB_COLOR_CHECK = "FF0000"

# FAST/Corality subtotal convention: a single rule above a subtotal, a double rule below
# a grand total (the "final answer" on a statement — Net Income, Total Assets, etc.).
_BORDER_SUBTOTAL = Border(top=Side(style="thin"))
_BORDER_GRAND_TOTAL = Border(top=Side(style="thin"), bottom=Side(style="double"))
FILL_SECTION_HEADER = PatternFill(start_color="FFF2F2F2", end_color="FFF2F2F2", fill_type="solid")


def style_total_row(ws: Worksheet, row: int, first_col_idx: int, last_col_idx: int,
                    grand: bool = False) -> None:
    """Apply the FAST/Corality subtotal/grand-total border convention across a row's data
    columns. `grand=True` adds the double-rule bottom edge reserved for a statement's
    final answer (Net Income, Total Assets, Closing Cash); everything else gets the
    single top rule that marks an ordinary subtotal."""
    border = _BORDER_GRAND_TOTAL if grand else _BORDER_SUBTOTAL
    for col_idx in range(first_col_idx, last_col_idx + 1):
        ws.cell(row=row, column=col_idx).border = border


def style_section_header_row(ws: Worksheet, row: int, first_col_idx: int, last_col_idx: int) -> None:
    """Light grey band across a section header's data columns, matching the label cell's
    existing bold/underline styling in column A."""
    for col_idx in range(first_col_idx, last_col_idx + 1):
        ws.cell(row=row, column=col_idx).fill = FILL_SECTION_HEADER


def new_workbook(template_path: str | None = None) -> Workbook:
    """Without a template, a blank workbook. With one, the template's .xlsm is reopened
    with its VBA project intact and every sheet dropped, so the builders repopulate a
    workbook that still carries the macros. openpyxl copies vbaProject.bin verbatim —
    verified to survive a full delete-and-rebuild of all sheets."""
    if template_path:
        wb = openpyxl.load_workbook(template_path, keep_vba=True)
        for name in list(wb.sheetnames):
            del wb[name]
        # Stale names would collide with the ones the builders re-register.
        for name in list(wb.defined_names):
            del wb.defined_names[name]
        return wb

    wb = openpyxl.Workbook()
    # Remove the default sheet; each builder module adds its own named sheet.
    wb.remove(wb.active)
    return wb


def col_letter(period_index: int) -> str:
    return openpyxl.utils.get_column_letter(FIRST_DATA_COL + period_index)


def build_workbook(inputs: ProjectInputs, timeline: Timeline, output_path: str,
                   template_path: str | None = None) -> Workbook:
    from cover import build_cover
    from assumptions_model import build_assumptions_model
    from assumptions_constant import build_assumptions_constant
    from assumptions_periodic_capex import build_assumptions_periodic_capex
    from assumptions_periodic_ops import build_assumptions_periodic_ops
    from calc_capex import build_calc_capex
    from calc_financing_cons import build_calc_financing_cons
    from calc_revenue_opex import build_calc_revenue_opex
    from calc_working_capital import build_calc_working_capital
    from calc_tax import build_calc_tax
    from calc_financing_ops import build_calc_financing_ops
    from calc_cfads import build_calc_cfads
    from fs_quarterly import build_fs_quarterly
    from fs_annual import build_fs_annual
    from valuation_selldown import build_valuation_selldown
    from batch_results import build_batch_results
    from stress_test import build_stress_test
    from check_control import build_check_control
    from dashboard import build_dashboard
    from legend import build_legend


    n_quarters = len(timeline.operations_quarters)
    tenor_quarters = min(inputs.financing.debt_tenor_years * 4, n_quarters)

    wb = new_workbook(template_path)
    build_cover(wb, len(timeline.construction_months),
                tenor_end_col=col_letter(tenor_quarters - 1),
                n_operating_quarters=n_quarters)
    build_assumptions_model(wb, timeline, inputs)
    build_assumptions_constant(wb, inputs)
    build_assumptions_periodic_capex(wb, timeline, inputs)
    build_assumptions_periodic_ops(wb, timeline, inputs)
    build_calc_capex(wb, timeline, inputs)
    build_calc_financing_cons(wb, timeline, inputs)
    build_calc_revenue_opex(wb, timeline, inputs)
    build_calc_working_capital(wb, timeline, inputs)
    build_calc_tax(wb, timeline, inputs)
    build_calc_financing_ops(wb, timeline, inputs)
    build_calc_cfads(wb, timeline, inputs)
    build_fs_quarterly(wb, timeline, inputs)
    build_fs_annual(wb, timeline, inputs)
    build_valuation_selldown(wb, timeline)
    build_batch_results(wb)
    build_stress_test(wb)
    build_check_control(wb, timeline)
    build_dashboard(wb, timeline)  # must be built last, once everything it links to exists
    build_legend(wb)  # index 0: opens here — a read-only "how to read this model" tab
    wb.save(output_path)
    return wb
