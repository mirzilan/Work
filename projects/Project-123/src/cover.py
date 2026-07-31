from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.workbook import Workbook

import assumptions_model as model
from workbook_builder import COLOR_INPUT, COLOR_LINK, TAB_COLOR_INPUT

CELL_CIRC_TOLERANCE = "B4"
CELL_MAX_ITERATIONS = "B5"
CELL_DEBT_SIZING_TOLERANCE = "B6"

CELL_MASTER_CHECK_LINK = "B9"

CELL_DRAWDOWN_METHOD = "B14"
CELL_DEBT_SIZING_MODE = "B17"
CELL_ACTIVE_SCENARIO = "B20"
CELL_DSRA_METHOD = "B23"

ABS_DSRA_METHOD = "$B$23"

DRAWDOWN_METHODS = ["Debt First", "Equity First", "Pari Passu"]
DEBT_SIZING_MODES = ["Fixed Gearing", "DSCR Sculpted"]
DSRA_METHODS = ["Cash Funded", "LC-Backed"]
N_SCENARIOS = 10

ROW_FRESHNESS_HEADER = 27
ROW_LAST_SOLVED = 28
ROW_SOLVE_STATUS = 29
ROW_SNAPSHOT_TABLE_HEADER = 31
ROW_FIRST_SNAPSHOT = 32

# (label, live-value formula) — every input a solve depends on. Tracking them individually
# rather than as one hashed checksum means the model names the assumption that moved.
TRACKED_INPUTS = [
    ("Total Capex", f"=Assumptions_Model!$B${model.ROW_TOTAL_CAPEX}"),
    ("Gearing", f"=Assumptions_Model!$B${model.ROW_DEBT_PCT}"),
    ("Interest Rate", f"=Assumptions_Model!$B${model.ROW_INTEREST_RATE}"),
    ("Debt Tenor (Years)", f"=Assumptions_Model!$B${model.ROW_DEBT_TENOR_YEARS}"),
    ("Target DSCR", f"=Assumptions_Model!$B${model.ROW_TARGET_DSCR}"),
    ("Annual Revenue", f"=Assumptions_Model!$B${model.ROW_ANNUAL_REVENUE}"),
    ("Opex % of Revenue", f"=Assumptions_Model!$B${model.ROW_OPEX_PCT}"),
    ("Tax Rate", f"=Assumptions_Model!$B${model.ROW_TAX_RATE}"),
    ("Useful Life (Years)", f"=Assumptions_Model!$B${model.ROW_USEFUL_LIFE_YEARS}"),
    ("Max Gearing", f"=Assumptions_Model!$B${model.ROW_MAX_GEARING}"),
    ("Routine Maint Capex %", f"=Assumptions_Model!$B${model.ROW_ROUTINE_MAINT_PCT}"),
    ("DSRA LC Fee", f"=Assumptions_Model!$B${model.ROW_DSRA_LC_FEE}"),
    ("Drawdown Method", f"={CELL_DRAWDOWN_METHOD.replace('B', '$B$')}"),
    ("Debt Sizing Mode", f"={CELL_DEBT_SIZING_MODE.replace('B', '$B$')}"),
    ("DSRA Funding Method", f"={ABS_DSRA_METHOD}"),
    ("Active Scenario", f"={CELL_ACTIVE_SCENARIO.replace('B', '$B$')}"),
    ("Capex Phasing (signature)", None),  # filled in at build time — needs the timeline width
]

# Row of the single flag Check_Control pulls. Derived, so adding a tracked input above
# cannot leave Check_Control pointing at a blank cell.
ROW_FRESHNESS_FLAG = ROW_FIRST_SNAPSHOT + len(TRACKED_INPUTS) + 1


def build_cover(wb: Workbook, n_construction_months: int = 24) -> Worksheet:
    ws = wb.create_sheet("Cover", 0)  # index 0: first sheet, opens here
    ws.sheet_properties.tabColor = TAB_COLOR_INPUT

    ws["A1"] = "Project 123 — Cover"
    ws["A1"].font = Font(bold=True, size=14)

    ws["A3"] = "VBA Solve Settings (used from Stage 1b onward)"
    ws["A3"].font = Font(bold=True)

    ws["A4"] = "Circularity Tolerance ($)"
    ws[CELL_CIRC_TOLERANCE] = 1.0
    ws[CELL_CIRC_TOLERANCE].font = Font(color=COLOR_INPUT)

    ws["A5"] = "Max Iterations (per convergence loop)"
    ws[CELL_MAX_ITERATIONS] = 100
    ws[CELL_MAX_ITERATIONS].font = Font(color=COLOR_INPUT)

    ws["A6"] = "Debt Sizing Tolerance ($, outer root-find on debt size D)"
    ws[CELL_DEBT_SIZING_TOLERANCE] = 100.0
    ws[CELL_DEBT_SIZING_TOLERANCE].font = Font(color=COLOR_INPUT)

    ws["A8"] = "Model Status"
    ws["A8"].font = Font(bold=True)
    status_cell = ws[CELL_MASTER_CHECK_LINK]
    status_cell.value = "=Check_Control!B3"
    status_cell.font = Font(bold=True, size=14)

    ws["A11"] = "Buttons (Stage 1b+): Solve Construction IDC | Solve Debt Sculpting | Goal Seek -> EIRR | Goal Seek -> PIRR"
    ws["A11"].font = Font(italic=True)
    ws["A12"] = "Placeholder rows only — Form Control buttons + macro assignment are a manual, one-time step (see VBA hand-off docs)."
    ws["A12"].font = Font(italic=True, size=9)

    ws["A14"] = "Construction Drawdown Method"
    ws["A14"].font = Font(bold=True)
    method_cell = ws[CELL_DRAWDOWN_METHOD]
    method_cell.value = DRAWDOWN_METHODS[2]
    method_cell.font = Font(color=COLOR_INPUT)

    validation = DataValidation(
        type="list",
        formula1=f'"{",".join(DRAWDOWN_METHODS)}"',
        allow_blank=False,
        showDropDown=False,
    )
    ws.add_data_validation(validation)
    validation.add(method_cell)

    ws["A15"] = (
        "Debt First: draw debt until facility exhausted, then equity. "
        "Equity First: draw equity until commitment exhausted, then debt. "
        "Pari Passu: draw both proportionally each month."
    )
    ws["A15"].font = Font(italic=True, size=9)

    ws["A17"] = "Debt Sizing Mode"
    ws["A17"].font = Font(bold=True)
    sizing_cell = ws[CELL_DEBT_SIZING_MODE]
    sizing_cell.value = DEBT_SIZING_MODES[1]
    sizing_cell.font = Font(color=COLOR_INPUT)

    sizing_validation = DataValidation(
        type="list",
        formula1=f'"{",".join(DEBT_SIZING_MODES)}"',
        allow_blank=False,
        showDropDown=False,
    )
    ws.add_data_validation(sizing_validation)
    sizing_validation.add(sizing_cell)

    ws["A18"] = (
        "Fixed Gearing: debt = gearing x Total Project Cost, level (PMT) amortisation. "
        "DSCR Sculpted: repayment locked to Target DSCR, debt size solved so the balance "
        "amortises to zero exactly at tenor end."
    )
    ws["A18"].font = Font(italic=True, size=9)

    ws["A20"] = "Active Scenario (1-10)"
    ws["A20"].font = Font(bold=True)
    scenario_cell = ws[CELL_ACTIVE_SCENARIO]
    scenario_cell.value = 1
    scenario_cell.font = Font(color=COLOR_INPUT, bold=True)

    scenario_validation = DataValidation(
        type="whole",
        operator="between",
        formula1=1,
        formula2=N_SCENARIOS,
        allow_blank=False,
        showErrorMessage=True,
        errorTitle="Invalid scenario",
        error=f"Enter a whole number between 1 and {N_SCENARIOS}.",
    )
    ws.add_data_validation(scenario_validation)
    scenario_validation.add(scenario_cell)

    ws["C20"] = "=Assumptions_Constant!$B$4"
    ws["C20"].font = Font(color=COLOR_LINK, italic=True)

    ws["A21"] = (
        "Drives the Active column on Assumptions_Constant and the Active row on both "
        "Periodic sheets — one switch moves every scenario-varying input together."
    )
    ws["A21"].font = Font(italic=True, size=9)

    ws["A23"] = "DSRA Funding Method"
    ws["A23"].font = Font(bold=True)
    dsra_cell = ws[CELL_DSRA_METHOD]
    dsra_cell.value = DSRA_METHODS[0]
    dsra_cell.font = Font(color=COLOR_INPUT)

    dsra_validation = DataValidation(
        type="list",
        formula1=f'"{",".join(DSRA_METHODS)}"',
        allow_blank=False,
        showDropDown=False,
    )
    ws.add_data_validation(dsra_validation)
    dsra_validation.add(dsra_cell)

    ws["A24"] = (
        "Cash Funded: CFADS is trapped to hold the reserve at target, released as the "
        "requirement falls. LC-Backed: no cash trapped — a recurring LC fee is charged on "
        "the requirement instead, and the fee is a tax-deductible P&L cost."
    )
    ws["A24"].font = Font(italic=True, size=9)

    _build_freshness_block(ws, wb, n_construction_months)

    _add_named_range(wb, "Cover_CircTolerance", "Cover", CELL_CIRC_TOLERANCE)
    _add_named_range(wb, "Cover_DebtSizingMode", "Cover", CELL_DEBT_SIZING_MODE)
    _add_named_range(wb, "Cover_MaxIterations", "Cover", CELL_MAX_ITERATIONS)
    _add_named_range(wb, "Cover_DebtSizingTolerance", "Cover", CELL_DEBT_SIZING_TOLERANCE)
    _add_named_range(wb, "Cover_DrawdownMethod", "Cover", CELL_DRAWDOWN_METHOD)
    _add_named_range(wb, "Cover_DSRAMethod", "Cover", CELL_DSRA_METHOD)
    _add_named_range(wb, "ActiveScenario", "Cover", CELL_ACTIVE_SCENARIO)

    ws.column_dimensions["A"].width = 45

    return ws


def _build_freshness_block(ws: Worksheet, wb: Workbook, n_construction_months: int) -> None:
    """Staged values are numbers, not formulas — change an assumption without re-solving and
    the model shows stale figures that look entirely valid. This block snapshots every input
    a solve depends on and flags any drift since the last solve."""
    from openpyxl.utils import get_column_letter

    ws.cell(row=ROW_FRESHNESS_HEADER, column=1, value="Solve Freshness").font = Font(bold=True)

    ws.cell(row=ROW_LAST_SOLVED, column=1, value="Last Solved")
    stamp = ws.cell(row=ROW_LAST_SOLVED, column=2)
    stamp.value = "(never)"
    stamp.font = Font(color=COLOR_INPUT)

    ws.cell(row=ROW_SOLVE_STATUS, column=1, value="Solve Status")
    ws.cell(row=ROW_SOLVE_STATUS, column=1).font = Font(bold=True)

    last_row = ROW_FIRST_SNAPSHOT + len(TRACKED_INPUTS) - 1
    status = ws.cell(row=ROW_SOLVE_STATUS, column=2)
    status.value = (
        f'=IF(COUNTIF(D{ROW_FIRST_SNAPSHOT}:D{last_row},0)=0,'
        f'"SOLVED - current","RE-RUN SOLVE - assumptions changed")'
    )
    status.font = Font(bold=True, size=12)

    for col, header in ((1, "Tracked Input"), (2, "Live"), (3, "At Last Solve"), (4, "Match")):
        cell = ws.cell(row=ROW_SNAPSHOT_TABLE_HEADER, column=col, value=header)
        cell.font = Font(bold=True)

    phasing_last_col = get_column_letter(3 + n_construction_months - 1)
    phasing_row = model.ROW_PHASING_PCT
    phasing_range = f"Assumptions_Model!$C${phasing_row}:${phasing_last_col}${phasing_row}"
    phasing_signature = f"=SUMPRODUCT({phasing_range},COLUMN({phasing_range}))"

    for i, (label, formula) in enumerate(TRACKED_INPUTS):
        row = ROW_FIRST_SNAPSHOT + i
        ws.cell(row=row, column=1, value=label)

        live = ws.cell(row=row, column=2)
        live.value = phasing_signature if formula is None else formula
        live.font = Font(color=COLOR_LINK)

        stored = ws.cell(row=row, column=3)
        stored.value = 0
        stored.font = Font(color=COLOR_INPUT)

        match = ws.cell(row=row, column=4)
        match.value = f"=IF(B{row}=C{row},1,0)"

    # Single flag Check_Control can pull: 1 when every tracked input still matches its
    # value at the last solve.
    freshness_flag = ws.cell(row=ROW_FRESHNESS_FLAG, column=4)
    freshness_flag.value = f"=IF(COUNTIF(D{ROW_FIRST_SNAPSHOT}:D{last_row},0)=0,1,0)"
    ws.cell(row=ROW_FRESHNESS_FLAG, column=1, value="Check: Solve is current")

    _add_named_range(wb, "LastSolvedStamp", "Cover", f"B{ROW_LAST_SOLVED}")
    _add_named_range(wb, "SolveStatus", "Cover", f"B{ROW_SOLVE_STATUS}")
    _add_named_range(wb, "SolveFreshnessFlag", "Cover", f"D{ROW_FRESHNESS_FLAG}")

    from openpyxl.workbook.defined_name import DefinedName

    wb.defined_names.add(
        DefinedName("SnapshotLive", attr_text=f"'Cover'!$B${ROW_FIRST_SNAPSHOT}:$B${last_row}")
    )
    wb.defined_names.add(
        DefinedName("SnapshotStored", attr_text=f"'Cover'!$C${ROW_FIRST_SNAPSHOT}:$C${last_row}")
    )
    wb.defined_names.add(
        DefinedName("SnapshotMatchFlags", attr_text=f"'Cover'!$D${ROW_FIRST_SNAPSHOT}:$D${last_row}")
    )


def _add_named_range(wb: Workbook, name: str, sheet: str, cell: str) -> None:
    import re

    from openpyxl.workbook.defined_name import DefinedName

    match = re.match(r"([A-Za-z]+)(\d+)", cell)
    col, row = match.group(1), match.group(2)
    wb.defined_names.add(DefinedName(name, attr_text=f"'{sheet}'!${col}${row}"))
