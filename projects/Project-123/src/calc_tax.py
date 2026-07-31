from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.workbook import Workbook

import assumptions_model as model
import calc_financing_ops as fin_ops
from inputs import ProjectInputs
from timeline import Timeline
# Import the source sheet's row constants rather than restating them: indexing another
# sheet with this module's own row numbers is how Calc_Tax once pulled Revenue while
# labelling it EBITDA.
from calc_revenue_opex import (
    ROW_EBITDA as REVOPEX_ROW_EBITDA,
    ROW_MAINT_TOTAL as REVOPEX_ROW_MAINT_TOTAL,
)
from workbook_builder import (
    FIRST_DATA_COL,
    COLOR_INPUT,
    COLOR_FORMULA,
    COLOR_LINK,
    TAB_COLOR_CALC,
    col_letter,
)

TAX_RATE_CELL = "B3"
USEFUL_LIFE_YEARS_CELL = "B4"
MAINT_LIFE_YEARS_CELL = "B5"

ABS_TAX_RATE = "$B$3"
ABS_USEFUL_LIFE = "$B$4"
ABS_MAINT_LIFE = "$B$5"

ROW_DATE_HEADER = 7
ROW_QUARTER_INDEX = 8

ROW_EBITDA = 10              # linked from Calc_Revenue_Opex
ROW_BASE_DEPRECIATION = 11   # straight-line on Total Project Cost, single vintage
ROW_MAINT_CAPEX = 12         # linked from Calc_Revenue_Opex
ROW_MAINT_DEPRECIATION = 13  # multi-vintage: every quarter's spend starts its own schedule
ROW_TOTAL_DEPRECIATION = 14
ROW_INTEREST = 15            # linked from Calc_Financing_Ops — the tax shield
ROW_EBT = 16
ROW_TAX = 17                 # max(EBT,0) * tax rate — no loss carryforward yet
ROW_NET_INCOME = 18

ROW_CHECK_HEADER = 21
ROW_CHECK_ACCUM_DEPR = 22
ROW_CHECK_ACCUM_MAINT_DEPR = 23


def build_calc_tax(wb: Workbook, timeline: Timeline, inputs: ProjectInputs) -> Worksheet:
    ws = wb.create_sheet("Calc_Tax")
    ws.sheet_properties.tabColor = TAB_COLOR_CALC

    ws["A1"] = "Calc_Tax — Quarterly Tax (base vintage + multi-vintage maintenance capex)"
    ws["A1"].font = Font(bold=True, size=12)

    ws["A3"] = "Tax Rate — linked from Assumptions_Model"
    ws[TAX_RATE_CELL] = f"=Assumptions_Model!$B${model.ROW_TAX_RATE}"
    ws[TAX_RATE_CELL].font = Font(color=COLOR_LINK)
    ws[TAX_RATE_CELL].number_format = "0.00%"

    ws["A4"] = "Useful Life (Years) — linked from Assumptions_Model"
    ws[USEFUL_LIFE_YEARS_CELL] = f"=Assumptions_Model!$B${model.ROW_USEFUL_LIFE_YEARS}"
    ws[USEFUL_LIFE_YEARS_CELL].font = Font(color=COLOR_LINK)

    # Structural, not scenario-varying: it sets how many columns each vintage's rolling
    # window spans, and those ranges are written out at build time. Editing this cell
    # rescales the charge without moving the window — change it in inputs.py and rebuild.
    ws["A5"] = "Maintenance Capex Useful Life (Years) — structural, rebuild to change"
    ws[MAINT_LIFE_YEARS_CELL] = inputs.reserves.maint_capex_useful_life_years
    ws[MAINT_LIFE_YEARS_CELL].font = Font(color=COLOR_INPUT)

    _label(ws, ROW_DATE_HEADER, "Period End Date")
    _label(ws, ROW_QUARTER_INDEX, "Operating Quarter #")
    _label(ws, ROW_EBITDA, "EBITDA ($) — linked from Calc_Revenue_Opex")
    _label(ws, ROW_BASE_DEPRECIATION, "Base Depreciation ($) — straight-line on Total Project Cost")
    _label(ws, ROW_MAINT_CAPEX, "Maintenance Capex ($) — linked from Calc_Revenue_Opex")
    _label(ws, ROW_MAINT_DEPRECIATION, "Maintenance Depreciation ($) — multi-vintage, straight-line")
    _label(ws, ROW_TOTAL_DEPRECIATION, "Total Depreciation ($)")
    _label(ws, ROW_INTEREST, "Interest Expense ($) — linked from Calc_Financing_Ops (tax shield)")
    _label(ws, ROW_EBT, "EBT ($) = EBITDA - Total Depreciation - Interest")
    _label(ws, ROW_TAX, "Tax ($) — no loss carryforward yet")
    _label(ws, ROW_NET_INCOME, "Net Income ($)")

    ws.cell(row=ROW_CHECK_HEADER, column=1, value="Checks").font = Font(bold=True)
    _label(ws, ROW_CHECK_ACCUM_DEPR, "Check: Accumulated Base Depreciation <= Total Project Cost")
    _label(ws, ROW_CHECK_ACCUM_MAINT_DEPR, "Check: Accumulated Maint Depreciation <= Cumulative Maint Capex")

    n_quarters = len(timeline.operations_quarters)
    last_cons_col = col_letter(len(timeline.construction_months) - 1)
    # Depreciable base is Total Project Cost (capex + capitalised IDC), not capex alone —
    # PP&E on the balance sheet opens at TPC, so depreciating capex only would strand the
    # IDC portion undepreciated for the life of the asset.
    tpc_ref = f"Calc_Capex!${last_cons_col}$17"
    quarterly_depr_expr = f"{tpc_ref}/({ABS_USEFUL_LIFE}*4)"

    useful_life_quarters = inputs.tax.useful_life_years * 4
    maint_life_quarters = inputs.reserves.maint_capex_useful_life_years * 4

    for i, period in enumerate(timeline.operations_quarters):
        col = col_letter(i)

        ws[f"{col}{ROW_DATE_HEADER}"] = period.end
        ws[f"{col}{ROW_DATE_HEADER}"].number_format = "mmm-yy"
        ws[f"{col}{ROW_QUARTER_INDEX}"] = i + 1

        _link(ws, col, ROW_EBITDA, f"Calc_Revenue_Opex!{col}{REVOPEX_ROW_EBITDA}")

        # Depreciation stops once fully depreciated (quarter index < useful_life_years*4)
        depr_cell = ws[f"{col}{ROW_BASE_DEPRECIATION}"]
        depr_cell.value = f"={quarterly_depr_expr}" if i < useful_life_quarters else 0
        depr_cell.font = Font(color=COLOR_FORMULA)
        depr_cell.number_format = "#,##0"

        _link(ws, col, ROW_MAINT_CAPEX, f"Calc_Revenue_Opex!{col}{REVOPEX_ROW_MAINT_TOTAL}")

        # Every quarter's maintenance spend opens its own straight-line vintage. Because
        # all vintages share one life, the charge in quarter t is just the spend still
        # inside the window divided by that life — one row, no NxN vintage matrix.
        window_start = col_letter(max(0, i - maint_life_quarters + 1))
        maint_depr = ws[f"{col}{ROW_MAINT_DEPRECIATION}"]
        maint_depr.value = (
            f"=SUM({window_start}{ROW_MAINT_CAPEX}:{col}{ROW_MAINT_CAPEX})/({ABS_MAINT_LIFE}*4)"
        )
        maint_depr.font = Font(color=COLOR_FORMULA)
        maint_depr.number_format = "#,##0"

        _formula(ws, col, ROW_TOTAL_DEPRECIATION,
                 f"={col}{ROW_BASE_DEPRECIATION}+{col}{ROW_MAINT_DEPRECIATION}")

        _link(ws, col, ROW_INTEREST, f"Calc_Financing_Ops!{col}{fin_ops.ROW_INTEREST}")

        # The DSRA LC fee is deliberately absent from this deduction. Its requirement is
        # forward debt service, so deducting it would run LC fee -> tax -> CFADS -> debt
        # service -> requirement -> LC fee: a genuine cycle that neither staged cell breaks
        # (the balance recursion ties quarter t+1 back to t). Treating the fee as
        # non-deductible keeps the sheet acyclic and understates the LC option's benefit
        # rather than overstating it.
        _formula(ws, col, ROW_EBT,
                 f"={col}{ROW_EBITDA}-{col}{ROW_TOTAL_DEPRECIATION}-{col}{ROW_INTEREST}")
        _formula(ws, col, ROW_TAX, f"=MAX({col}{ROW_EBT},0)*{ABS_TAX_RATE}")
        _formula(ws, col, ROW_NET_INCOME, f"={col}{ROW_EBT}-{col}{ROW_TAX}")

    last_col = col_letter(n_quarters - 1)
    first_col = col_letter(0)

    _check(ws, last_col, ROW_CHECK_ACCUM_DEPR,
           f"=IF(SUM({first_col}{ROW_BASE_DEPRECIATION}:{last_col}{ROW_BASE_DEPRECIATION})"
           f"<={tpc_ref}+0.01,1,0)")

    # Vintages opened near the end of life run past the model horizon, so the charge is
    # truncated, never overstated — this catches the window being wired the wrong way.
    _check(ws, last_col, ROW_CHECK_ACCUM_MAINT_DEPR,
           f"=IF(SUM({first_col}{ROW_MAINT_DEPRECIATION}:{last_col}{ROW_MAINT_DEPRECIATION})"
           f"<=SUM({first_col}{ROW_MAINT_CAPEX}:{last_col}{ROW_MAINT_CAPEX})+0.01,1,0)")

    _add_named_range(wb, "Tax_LastCol", "Calc_Tax", f"{last_col}1")
    _add_named_range(wb, "Tax_AccumDeprCheck", "Calc_Tax", f"{last_col}{ROW_CHECK_ACCUM_DEPR}")
    _add_named_range(wb, "Tax_AccumMaintDeprCheck", "Calc_Tax", f"{last_col}{ROW_CHECK_ACCUM_MAINT_DEPR}")

    ws.freeze_panes = ws.cell(row=ROW_NET_INCOME + 1, column=FIRST_DATA_COL)
    ws.column_dimensions["A"].width = 56

    return ws


def _label(ws: Worksheet, row: int, label: str) -> None:
    ws.cell(row=row, column=1, value=label)


def _link(ws: Worksheet, col: str, row: int, formula: str) -> None:
    cell = ws[f"{col}{row}"]
    cell.value = f"={formula}"
    cell.font = Font(color=COLOR_LINK)
    cell.number_format = "#,##0"


def _formula(ws: Worksheet, col: str, row: int, formula: str) -> None:
    cell = ws[f"{col}{row}"]
    cell.value = formula
    cell.font = Font(color=COLOR_FORMULA)
    cell.number_format = "#,##0"


def _check(ws: Worksheet, col: str, row: int, formula: str) -> None:
    cell = ws[f"{col}{row}"]
    cell.value = formula
    cell.font = Font(color=COLOR_FORMULA)


def _add_named_range(wb: Workbook, name: str, sheet: str, cell: str) -> None:
    import re

    from openpyxl.workbook.defined_name import DefinedName

    match = re.match(r"([A-Za-z]+)(\d+)", cell)
    col, row = match.group(1), match.group(2)
    wb.defined_names.add(DefinedName(name, attr_text=f"'{sheet}'!${col}${row}"))
